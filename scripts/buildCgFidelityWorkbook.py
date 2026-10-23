#!/usr/bin/env python3
"""Build the local, reviewable Excel fidelity/CG register from project evidence.

Requires openpyxl >= 3.1 and Pillow. No remote media is fetched. Existing manual
review cells are preserved when the workbook is rebuilt. Project-generated art
is labelled as project art, never as an independently validated canon reference.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any

from PIL import Image, ImageOps
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.workbook.properties import CalcProperties
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs" / "cg"
PLAN = DOCS / "roster-production-2026-10-18.json"
BIBLE = DOCS / "roster-fidelity-bible-2026-10-18.json"
DEFAULT_OUTPUT = DOCS / "Multiverse_Breach_Fidelite_CG.xlsx"
GLOBAL_BATCH = DOCS / "roster-global-batch-2026-10-23.json"
GLOBAL_PROMPT = DOCS / "PROMPT_GLOBAL_ROSTER_1_1.txt"
LINK_BASE = DOCS
THUMBNAIL_CACHE = ROOT / ".tmp" / "cg-fidelity-workbook-thumbnails" / "v1"
BLUE = "172C4A"
CYAN = "DDF3F6"
LIGHT = "F1F5FA"
GREEN = "DFF1E5"
AMBER = "FFF0CC"
RED = "FBE3DF"
WHITE = "FFFFFF"
INK = "24344B"
MUTED = "64748B"
BODY_FONT = Font(name="Aptos", size=10, color=INK)
BODY_ALIGNMENT = Alignment(vertical="top", wrap_text=True)
LINK_FONT = Font(name="Aptos", size=10, color="156B9C", underline="single")
INTERNAL_CANON_VISUAL_ROLES = {
    "canonical-original-project-sprite-pack",
    "canonical-original-project-portrait",
}

AXES = [
    "Incarnation / continuité", "Silhouette / proportions", "Anatomie",
    "Visage / marques", "Tenue / palette", "Équipement",
    "Cohérence univers / époque", "Source visuelle indépendante",
]
CONTROL_HEADERS = ["ID personnage", "Personnage", "Univers"] + AXES + [
    "Axes conformes", "Objectif 1:1", "Décision manuelle", "Commentaire QA",
    "Nom du relecteur", "Date de revue", "Aller au roster",
    "Cible d'incarnation manuelle", "Référence visuelle validée (URL / fichier)",
]
MATRIX_HEADERS = [
    "ID personnage", "Personnage", "Univers", "Priorité", "CG", "Type",
    "Nature du thème", "Thème / contrainte", "Pose distincte prévue / livrée",
    "Décor distinct prévu / livré", "ID décor", "État réel de production",
    "Qualification du thème", "Incarnation retenue", "Source du thème",
    "Références personnage", "Image actuelle", "PNG natif / source",
    "Anciennes variantes", "Contrôle visuel", "Limites connues",
    "Revue manuelle", "Commentaire QA", "Aller au roster",
]
ROSTER_HEADERS = [
    "Visuel projet actuel", "ID personnage", "Personnage", "Univers", "Priorité",
    "Incarnation / continuité", "Qualification incarnation", "Silhouette",
    "Anatomie", "Visage", "Tenue", "Palette", "Équipement",
    "Marques distinctives", "Mélanges interdits", "Œuvre source", "Époque",
    "Lieux cohérents", "Architecture / matériaux", "Palette univers",
    "Accessoires univers", "Limites de cohérence", "Référence visuelle",
    "Revue source", "Fidélité 1:1", "CG livrées", "CG à produire", "CG N/A",
    "Fichier visuel", "Sources et statut", "À vérifier", "Matrice CG",
    "Contrôles", "Nature du visuel affiché",
]


def read_json(path: Path, fallback: Any = None) -> Any:
    if not path.is_file():
        return fallback
    return json.loads(path.read_text(encoding="utf-8"))


def plain(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Oui" if value else "Non"
    if isinstance(value, dict):
        return "\n".join(f"{k}: {plain(v)}" for k, v in value.items() if v not in (None, [], {}, ""))
    if isinstance(value, (list, tuple)):
        return "\n".join(plain(v) for v in value if v not in (None, [], {}, ""))
    return str(value)


def cell_text(value: Any) -> Any:
    if isinstance(value, (int, float, date, datetime)):
        return value
    value = plain(value)
    # Excel text limits and formula injection protection for source prose.
    value = re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F]", "", value)[:32760]
    if value.startswith(("=", "+", "-", "@")):
        value = "'" + value
    return value


def local_path(raw: str | None) -> Path | None:
    if not raw or raw.startswith(("https://", "http://")):
        return None
    clean = raw.split("#", 1)[0]
    p = Path(clean)
    if p.is_absolute() and p.is_file():
        return p
    if clean.startswith(("/cg/", "/sprites/", "/visuals/", "/images/")):
        p = ROOT / "public" / clean.lstrip("/")
    else:
        p = ROOT / clean.lstrip("/")
    return p if p.is_file() else None


def ref_lines(refs: Any) -> str:
    if not refs:
        return ""
    if not isinstance(refs, list):
        refs = [refs]
    lines = []
    for source in refs:
        if isinstance(source, dict):
            ref = source.get("url") or source.get("path") or source.get("ref") or ""
            role = source.get("role") or source.get("type") or "source"
            status = source.get("status") or "statut non documenté"
            lines.append(f"[{status} · {role}] {ref}")
        else:
            lines.append(str(source))
    return "\n".join(lines)


def collect_versions(plan: dict) -> tuple[dict, dict]:
    """Only register actual files; planned filenames cannot become delivered CG."""
    versions: dict[str, list[dict]] = defaultdict(list)
    entries = {entry["id"]: entry for entry in plan["entries"]}
    # The production plan may point at current revisions. Historical files must
    # come from the immutable published manifest, never from currentAssets.
    baseline = read_json(ROOT / "public/cg/cg-manifest.json", {}).get("entries", [])
    for original in baseline:
        entry = {**original, **entries.get(original["id"], {})}
        entries[original["id"]] = entry
        assets = original.get("assets") or {}
        master = (assets.get("master") or {}).get("path")
        if not local_path(master):
            continue
        versions[entry["id"]].append({
            "cgId": entry["id"], "revisionId": "openai-v1-historique",
            "imagePath": master,
            "sourcePath": (assets.get("source") or {}).get("path"),
            "thumbnailPath": (assets.get("thumbnail") or {}).get("path"),
            "publishedAt": original.get("publishedAt") or "2026-08-01",
            "backgroundId": entry.get("backgroundId"),
            "themeBrief": entry.get("themeBrief"),
            "poseBrief": "Version originale publiée; pose à comparer directement dans l’image historique.",
            "sourceRefs": original.get("sourceRefs") or "public/cg/cg-manifest.json",
            "visualReview": {"status": "historical-project-art-not-1-to-1-certified"},
            "assets": assets,
        })

    extra = read_json(ROOT / "src/game/cg/cgRosterArtwork.generated.json", {})
    for entry in extra.get("entries", []):
        entries[entry["id"]] = {**entries.get(entry["id"], {}), **entry}
        assets = entry.get("assets") or {}
        master = (assets.get("master") or {}).get("path") or entry.get("imagePath")
        if not local_path(master):
            continue
        candidate = {
            **entry, "cgId": entry["id"],
            "revisionId": entry.get("revisionId") or f"fidelity-{entry.get('publishedAt', 'undated')}",
            "imagePath": master,
            "sourcePath": (assets.get("source") or {}).get("path") or entry.get("sourcePath"),
            "thumbnailPath": (assets.get("thumbnail") or {}).get("path") or entry.get("thumbnailPath"),
            "themeBrief": entry.get("themeBrief") or entry.get("promptSummary", {}).get("fr"),
            "poseBrief": entry.get("poseBrief") or entry.get("poseDescription"),
        }
        # Entries added after the original production plan may have no base row.
        if not any(v.get("imagePath") == candidate["imagePath"] for v in versions[entry["id"]]):
            versions[entry["id"]].append(candidate)

    # Every base must precede its revisions, including new roster artwork.
    # Publication dates have day resolution; stable ordering on the same day
    # must therefore preserve the registry's explicit revision order.
    registry = read_json(ROOT / "src/game/cg/cgRevisions.generated.json", {})
    for revision in registry.get("revisions", []):
        if local_path(revision.get("imagePath")):
            versions[revision["cgId"]].append(revision)

    for cg_id, items in versions.items():
        # Stable order preserves publication order for same-date revisions.
        versions[cg_id] = sorted(items, key=lambda v: str(v.get("publishedAt", "")))
        # Keep the immutable historical base first regardless of later labels.
        versions[cg_id].sort(key=lambda v: v.get("revisionId") != "openai-v1-historique")
    return dict(versions), entries


def read_manual(output: Path) -> dict:
    result: dict[str, dict] = {"Contrôles": {}, "Matrice CG": {}}
    if not output.is_file():
        return result
    old = load_workbook(output, read_only=True, data_only=False)
    try:
        for sheet_name, key_names, fields in [
            ("Contrôles", ["ID personnage"], AXES + ["Décision manuelle", "Commentaire QA", "Nom du relecteur", "Date de revue", "Cible d'incarnation manuelle", "Référence visuelle validée (URL / fichier)"]),
            ("Matrice CG", ["ID personnage", "CG"], ["Revue manuelle", "Commentaire QA"]),
        ]:
            if sheet_name not in old:
                continue
            rows = old[sheet_name].iter_rows(min_row=3, values_only=True)
            headers = next(rows)
            index = {value: i for i, value in enumerate(headers)}
            if any(key not in index for key in key_names):
                continue
            for row in rows:
                key = tuple(row[index[name]] for name in key_names)
                if not key[0]:
                    continue
                values = {field: row[index[field]] for field in fields if field in index and index[field] < len(row)}
                if any(value not in (None, "", "À contrôler", "Non revue") for value in values.values()):
                    result[sheet_name][key] = values
    finally:
        old.close()
    return result


def base_sheet(wb: Workbook, title: str, headers: list[str], subtitle: str, freeze: str = "D4"):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:3"
    ws.freeze_panes = freeze
    ws.append([title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.cell(1, 1).font = Font(name="Aptos Display", size=21, bold=True, color=WHITE)
    ws.cell(1, 1).fill = PatternFill("solid", fgColor=BLUE)
    ws.row_dimensions[1].height = 34
    ws.append([subtitle])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    ws.cell(2, 1).font = Font(name="Aptos", size=10, italic=True, color=MUTED)
    ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 31
    ws.append(headers)
    ws.row_dimensions[3].height = 32
    for cell in ws[3]:
        cell.font = Font(name="Aptos", bold=True, color=WHITE, size=10)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    return ws


def finish_table(ws, name: str, widths: dict[int, float] | None = None, height: float = 58):
    last_col = get_column_letter(ws.max_column)
    if ws.max_row >= 4:
        table = Table(displayName=name, ref=f"A3:{last_col}{ws.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = (widths or {}).get(col, 28)
    for row in ws.iter_rows(min_row=4):
        ws.row_dimensions[row[0].row].height = height
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = BODY_ALIGNMENT
    ws.auto_filter.ref = f"A3:{last_col}{ws.max_row}"
    ws.print_options.horizontalCentered = True


def append_row(ws, values):
    ws.append([cell_text(v) for v in values])


def link_cell(cell, target: str | Path | None, label: str | None = None):
    if not target:
        return
    target = str(target)
    if label is not None:
        cell.value = label
    if target.startswith("#"):
        cell.hyperlink = target
    elif target.startswith(("https://", "http://")):
        cell.hyperlink = target
    else:
        p = local_path(target)
        if p:
            cell.hyperlink = Path(os.path.relpath(p, LINK_BASE)).as_posix()
    if cell.hyperlink:
        cell.font = LINK_FONT


def image_bytes(path: Path, size: tuple[int, int]) -> BytesIO:
    stat = path.stat()
    fingerprint = f"{path.resolve()}|{stat.st_size}|{stat.st_mtime_ns}|{size[0]}x{size[1]}|png-white-v1"
    key = hashlib.sha256(fingerprint.encode("utf-8")).hexdigest()
    cache = THUMBNAIL_CACHE / f"{key}.png"
    if cache.is_file():
        return BytesIO(cache.read_bytes())
    with Image.open(path) as source:
        im = ImageOps.exif_transpose(source).convert("RGBA")
        im.thumbnail(size, Image.Resampling.LANCZOS)
        canvas = Image.new("RGB", size, "#f4f6fa")
        canvas.paste(im, ((size[0] - im.width) // 2, (size[1] - im.height) // 2), im)
        stream = BytesIO()
        canvas.save(stream, format="PNG", optimize=True)
        THUMBNAIL_CACHE.mkdir(parents=True, exist_ok=True)
        cache.write_bytes(stream.getvalue())
        stream.seek(0)
        return stream


def embed(ws, anchor: str, path: Path | None, streams: list, size=(112, 112)) -> bool:
    if not path:
        return False
    try:
        stream = image_bytes(path, size)
        image = XLImage(stream)
        image.width, image.height = size
        image.anchor = anchor
        ws.add_image(image)
        streams.append(stream)
        return True
    except (OSError, ValueError):
        return False


def add_list_validation(ws, col: int, list_col: int, last_list_row: int):
    name = {1: "FidelityAxisStates", 2: "FidelityDecisions", 3: "CgReviewStates"}[list_col]
    validation = DataValidation(type="list", formula1=name)
    validation.errorTitle = "Choisir un état de la liste"
    validation.error = "Les états doivent permettre une revue cohérente."
    validation.showErrorMessage = True
    validation.errorStyle = "stop"
    validation.allow_blank = True
    ws.add_data_validation(validation)
    validation.add(f"{get_column_letter(col)}4:{get_column_letter(col)}{ws.max_row}")


def subject_visual(subject: dict, bible: dict, current: dict, entry_by_id: dict) -> tuple[str, str]:
    for candidate in bible.get("localVisuals", []):
        if isinstance(candidate, dict) and candidate.get("selectedInRuntime") and local_path(candidate.get("path")):
            return candidate["path"], "Sprite / planche réellement sélectionné dans le jeu — pas une source canon indépendante"
    for cg_id in subject.get("entryIds", []):
        entry = entry_by_id.get(cg_id, {})
        if entry.get("type") == "characterSolo" and cg_id in current:
            value = current[cg_id]
            return value.get("thumbnailPath") or value.get("imagePath"), "CG actuelle générée du projet — pas une source canon indépendante"
    for candidate in bible.get("localVisuals", []) + subject.get("identity", {}).get("localCandidates", []):
        path = candidate.get("path") if isinstance(candidate, dict) else candidate
        if local_path(path):
            return path, "Planche / sprite actuel du projet — pas une source canon indépendante"
    return "", "Aucune image locale disponible"


def theme_for(bible: dict, slot: str, type_name: str) -> dict:
    themes = (bible.get("cgRequirements") or {}).get("themes") or {}
    if isinstance(themes, dict):
        value = themes.get(slot) or themes.get(type_name) or {}
        return value if isinstance(value, dict) else {"themeBrief": value}
    for value in themes:
        if isinstance(value, dict) and (value.get("slot") == slot or value.get("type") == type_name):
            return value
    return {}


def compact_review(review: dict | str | None) -> str:
    if not isinstance(review, dict):
        return plain(review) or "À contrôler"
    text = [f"État: {review.get('status', 'non renseigné')}", "Fidélité canon 1:1 non certifiée"]
    for key, label in [("poseDistinct", "Pose distincte"), ("backgroundDistinct", "Décor distinct")]:
        if key in review:
            value = "oui" if review[key] is True else "non" if review[key] is False else plain(review[key])
            text.append(f"{label}: {value}")
    if review.get("reviewedAt"):
        text.append("Revue documentée: " + plain(review["reviewedAt"]))
    independent = review.get("independentReviewFile") or review.get("independentReview")
    if independent:
        text.append("Reçu de comparaison indépendante: " + plain(independent))
    return "\n".join(text)


def collect_visual_evidence(bible: dict) -> dict:
    """Source bitmaps with provenance; project art needs explicit OC authority."""
    result = defaultdict(list)
    seen = {}
    scoped_reuses = []

    def add(char_id, item, base=ROOT, receipt="", resolve_reuse=False):
        if not char_id or not isinstance(item, dict):
            return
        if item.get("excludedFromGeneration") or item.get("excluded"):
            return
        raw = item.get("path") or item.get("localPath") or item.get("file")
        if not raw:
            return
        candidate = local_path(str(raw)) or (base / str(raw))
        if not candidate.is_file() or candidate.suffix.lower() not in {".png", ".jpg", ".jpeg", ".webp"}:
            return
        role = str(item.get("role") or item.get("kind") or "")
        # Diagnostic enlargements stay attached to their parent source. They
        # must never inflate the number of independent references or identities.
        if role == "derived-source-inspection-crop" or re.search(r"derived.*inspection|inspection.*crop", role, re.I):
            return
        incarnation = (bible.get(char_id, {}).get("incarnation") or {})
        original_scoped_environment = (
            role == "original-project-environment-reference"
            and item.get("sourceScope") == "original-project-canon"
            and item.get("referencePurpose") == "environment-only"
            and item.get("viewed") is True
            and (item.get("usedForIdentity") is False or item.get("identityUse") is False)
        )
        original_authority = (
            role in INTERNAL_CANON_VISUAL_ROLES or original_scoped_environment
        ) and isinstance(incarnation, dict) and incarnation.get("status") == "original-project-canon"
        # Moving a saved project bitmap under docs/ does not turn it into an
        # external source. Explicit source copies and old comparison art stay
        # out of the source column unless they are the qualified OC authority.
        if not original_authority and (
            item.get("usedAsIndependentReference") is False
            or item.get("referencePurpose") in {"legacy-runtime-comparison-only", "historical-comparison-only"}
            or role.startswith("comparison-only")
            or role == "same-source-copy"
        ):
            return
        purpose = item.get("referencePurpose")
        if item.get("referenceScope") == "iconic-outfit-only" or role == "independent-iconic-outfit-reference":
            purpose = "costume-only"
        elif not purpose and (role == "independent-equipment-reference" or "weapon-reference" in role or "weapon-reload-reference" in role):
            purpose = "equipment-only"
        elif not purpose and role == "independent-environment-reference":
            purpose = "environment-only"
        purpose_slot = {"costume-only": "CG13", "equipment-only": "CG02", "environment-only": "CG03"}.get(purpose)
        scoped = purpose in {"costume-only", "equipment-only", "environment-only", "body-costume-support"}
        # A single OC portrait can document its own equipment or environment in
        # separate slots. Resolve these uses after identity authorities so file
        # traversal order cannot turn the same project bitmap into an external
        # source or discard its expressly limited use.
        if not original_authority and scoped and isinstance(incarnation, dict) and incarnation.get("status") == "original-project-canon" and item.get("usedForIdentity") is False and not resolve_reuse:
            scoped_reuses.append((char_id, item, base, receipt))
            return
        reused_authority = None
        if resolve_reuse:
            candidate_hash = hashlib.sha256(candidate.read_bytes()).hexdigest()
            for authority in result[char_id]:
                if authority.get("sourceScope") == "original-project-canon" and not authority.get("reusedAuthorityBitmap") and hashlib.sha256(Path(authority["path"]).read_bytes()).hexdigest() == candidate_hash:
                    reused_authority = authority
                    original_authority = True
                    break
        project_location = str(raw).startswith(("/sprites/", "/cg/", "public/sprites/", "public/cg/")) or "/generated_images/" in str(candidate)
        if not original_authority and (project_location or "generated" in role or "project" in role):
            return
        expected_hash = item.get("sha256")
        if expected_hash and hashlib.sha256(candidate.read_bytes()).hexdigest() != expected_hash:
            raise ValueError(f"Independent reference hash mismatch: {candidate}")
        # Retain the established bitmap de-duplication everywhere else. Only a
        # positively matched internal authority has separately documented
        # same-bitmap uses; an unrelated repeated review cannot add identity or
        # costume rows for an equipment-only source.
        key = (char_id, str(candidate.resolve()))
        if reused_authority:
            key += (purpose, item.get("slot") or purpose_slot)
        if key in seen:
            if item.get("viewed") is True:
                seen[key]["viewed"] = True
                seen[key]["receipt"] = receipt or seen[key].get("receipt", "")
            return
        reference = {**item, "path": str(candidate.resolve()), "receipt": receipt, "sourceScope": "original-project-canon" if original_authority else "external-reference", "referencePurpose": purpose, "slot": item.get("slot") or purpose_slot}
        if reused_authority:
            reference.update({"reusedAuthorityBitmap": True, "authorityRole": reused_authority.get("role"), "authorityPath": reused_authority["path"]})
        seen[key] = reference
        result[char_id].append(reference)

    curated = DOCS / "workbook-visual-comparisons-2026-10-18.json"
    for item in read_json(curated, {}).get("references", []):
        add(item.get("characterId"), item, receipt=str(curated.relative_to(ROOT)))
    for char_id, dossier in bible.items():
        for visual in dossier.get("localVisuals", []):
            if isinstance(visual, dict):
                add(char_id, visual, receipt=str(BIBLE.relative_to(ROOT)))
    review_files = set((DOCS / "reviews").rglob("*.json"))
    review_files.update(DOCS.glob("*reference-evidence-*.json"))
    review_files.update((DOCS / "references").rglob("*receipt*.json"))
    for path in sorted(review_files):
        review = read_json(path, {})
        if not isinstance(review, dict):
            continue
        review_character = review.get("heroId") or review.get("characterId")
        for ref in review.get("references", []):
            if isinstance(ref, dict):
                add(review_character, {**ref, "limitations": ref.get("limitations") or review.get("limitations") or review.get("limits")}, receipt=str(path.relative_to(ROOT)))
        # A documented enlargement of an OC's own sprite can make the source
        # readable in Excel. It stays attached to the original pack reference,
        # never becomes a second reference or an external franchise source.
        crop = review.get("derivedInspection") or {}
        parents = [ref for ref in review.get("references", []) if isinstance(ref, dict) and ref.get("role") == "canonical-original-project-sprite-pack"]
        if isinstance(crop, dict) and crop.get("role") == "nearest-neighbor-inspection-crop-of-selected-sprite" and crop.get("viewed") is True and crop.get("sha256") and len(parents) == 1:
            crop_path = local_path(crop.get("path"))
            parent_path = local_path(parents[0].get("path"))
            if crop_path and crop_path.is_file() and parent_path:
                if hashlib.sha256(crop_path.read_bytes()).hexdigest() != crop["sha256"]:
                    raise ValueError(f"Internal source inspection crop hash mismatch: {crop_path}")
                for reference in result.get(review_character, []):
                    if reference.get("sourceScope") == "original-project-canon" and reference.get("path") == str(parent_path.resolve()):
                        reference["inspectionPath"] = str(crop_path.resolve())
                        reference["inspectionReceipt"] = str(path.relative_to(ROOT))
        # Tiny characters in publisher screenshots can also use a documented
        # diagnostic crop. Keep the source row/link on the original bitmap and
        # verify exact crop pixels instead of treating the zoom as a new source.
        for detail in crop if isinstance(crop, list) else []:
            if not isinstance(detail, dict) or detail.get("role") != "nearest-neighbor-inspection-crop-of-official-source" or detail.get("viewed") is not True or detail.get("notIndependentReference") is not True:
                continue
            detail_path, parent_path = local_path(detail.get("path")), local_path(detail.get("parent"))
            operation = detail.get("operation") or {}
            rectangle, scale = operation.get("cropXYWH"), operation.get("nearestScale")
            if not detail_path or not parent_path or not isinstance(rectangle, list) or len(rectangle) != 4 or not isinstance(scale, int) or scale < 1:
                continue
            if hashlib.sha256(detail_path.read_bytes()).hexdigest() != detail.get("sha256") or hashlib.sha256(parent_path.read_bytes()).hexdigest() != detail.get("parentSha256"):
                raise ValueError(f"External source inspection provenance mismatch: {detail_path}")
            x, y, width, height = rectangle
            with Image.open(parent_path) as parent_image, Image.open(detail_path) as detail_image:
                if x < 0 or y < 0 or width < 1 or height < 1 or x + width > parent_image.width or y + height > parent_image.height:
                    raise ValueError(f"External source inspection bounds mismatch: {detail_path}")
                expected = parent_image.convert("RGB").crop((x, y, x + width, y + height)).resize((width * scale, height * scale), Image.Resampling.NEAREST)
                if detail_image.size != expected.size or detail_image.convert("RGB").tobytes() != expected.tobytes():
                    raise ValueError(f"External source inspection pixels mismatch: {detail_path}")
            for reference in result.get(review_character, []):
                if reference.get("sourceScope") == "external-reference" and reference.get("path") == str(parent_path.resolve()):
                    reference["inspectionPath"] = str(detail_path.resolve())
                    reference["inspectionReceipt"] = str(path.relative_to(ROOT))
    universal = ROOT / ".tmp/fidelity-2026-10-18/minions/universal-reference-metadata.json"
    for char_id in ["bob_minions", "kevin_minions", "stuart_minions"]:
        for ref in read_json(universal, []):
            add(char_id, ref, universal.parent, str(universal.relative_to(ROOT)))
    mistral_meta = ROOT / ".tmp/cg-complete-2026-10-18/mistral/render-reference-metadata.json"
    add("mistral_mgr", read_json(mistral_meta, {}), receipt=str(mistral_meta.relative_to(ROOT)))
    kelly_doc = ROOT / "docs/audits/fidelity-spider-kelly-source-review-2026-10-18.json"
    kelly = read_json(kelly_doc, {})
    for filename, ref in zip(["body-cover.jpg", "body-gameplay.jpg"], kelly.get("references", [])):
        add("cyber_spider_kelly", {**ref, "path": str(ROOT / ".tmp/fidelity-spider-kelly-2026-10-18/references" / filename), "limitations": kelly.get("unresolvedVisualDetails")}, receipt=str(kelly_doc.relative_to(ROOT)))
    for char_id, item, base, receipt in scoped_reuses:
        add(char_id, item, base, receipt, resolve_reuse=True)
    return dict(result)


def collect_production_notes() -> dict:
    notes = {}
    def cg_slot(value):
        if isinstance(value, dict):
            value = value.get("slot") or value.get("slotNumber")
        digits = str(value).removeprefix("CG")
        return f"CG{int(digits):02}" if digits.isdigit() and 1 <= int(digits) <= 15 else None

    receipts = set(DOCS.glob("*complete-*.json")) | set(DOCS.glob("*wave[0-9]-*.json"))
    for path in sorted(receipts):
        receipt = read_json(path, {})
        entries = receipt.get("entries") or []
        first_entry = entries[0] if entries and isinstance(entries[0], dict) else {}
        char_id = receipt.get("characterId") or receipt.get("heroId") or first_entry.get("characterId") or first_entry.get("heroId")
        if not char_id:
            continue
        # A genuine all-refused CG batch has no delivered entry. Admit its
        # explicit fifteen-slot character plan without misclassifying sprite
        # refusals or source-only reports as CG attempts.
        has_cg_entries = any(isinstance(entry, dict) and str(entry.get("id", "")).startswith("cg:") for entry in entries)
        linked_plan = read_json(ROOT / str(receipt.get("planFile") or ""), {})
        planned_slots = linked_plan.get("slots") or []
        cg_types = {
            "characterSolo", "weaponSolo", "decorSolo", "coherentScene",
            "actionScene", "introPose", "victoryPose", "defeatPose",
            "beachFamily", "maidService", "goofy", "alignmentSwap",
            "genderSwap", "iconicOutfitSwap", "zombieVersion",
            "firstStep", "futureExperienced",
        }
        explicit_empty_cg_batch = (
            not entries and receipt.get("plannedSlots") == 15
            and (linked_plan.get("characterId") or linked_plan.get("heroId")) == char_id
            and isinstance(planned_slots, list) and len(planned_slots) == 15
            and {cg_slot(row) for row in planned_slots} == {f"CG{i:02}" for i in range(1, 16)}
            and all(isinstance(row, dict) and row.get("type") in cg_types for row in planned_slots)
        )
        if not has_cg_entries and not explicit_empty_cg_batch:
            continue
        for blockage in [*receipt.get("blockages", []), *receipt.get("toolFailures", []), *receipt.get("blockedSlots", [])]:
            if not isinstance(blockage, dict):
                continue
            slot_id = cg_slot(blockage)
            if not slot_id:
                continue
            categories = set()
            for attempt in [blockage, *blockage.get("attempts", [])]:
                values = attempt.get("categories") or []
                if isinstance(values, str):
                    values = [values]
                for value in [attempt.get("category"), *values]:
                    if value:
                        categories.add(str(value))
            category = ", ".join(sorted(categories)) or "non renseignée"
            notes[(char_id, slot_id)] = {"blocked": True, "text": f"Refus du service de génération; catégorie communiquée: {category}. Une ancienne image éventuelle reste conservée; aucune nouvelle révision n'est déclarée livrée. Reçu: " + str(path.relative_to(ROOT))}
        production_plan = receipt.get("productionPlan") or []
        if isinstance(production_plan, dict):
            production_plan = production_plan.get("slots") or []
        if not isinstance(production_plan, list):
            production_plan = []
        unattempted = [*receipt.get("unattemptedSlots", []), *(entry.get("slot") for entry in production_plan if isinstance(entry, dict) and str(entry.get("status", "")).startswith("not-attempted"))]
        for slot in unattempted:
            slot_id = cg_slot(slot)
            if not slot_id:
                continue
            notes.setdefault((char_id, slot_id), {"blocked": False, "text": "Ce créneau n'a pas encore été tenté dans le lot documenté. Reçu: " + str(path.relative_to(ROOT))})
    return notes


def collect_global_batch(plan: dict, current: dict, manifest_path: Path | None = None,
                         prompt_path: Path | None = None) -> dict:
    """Link the common instruction only when its whole-roster snapshot is current.

    Queue disposition is preparation metadata. It must never increase the
    delivered image count or certify fidelity on behalf of a visual reviewer.
    """
    manifest_path = manifest_path or GLOBAL_BATCH
    prompt_path = prompt_path or GLOBAL_PROMPT
    if not manifest_path.is_file() and not prompt_path.is_file():
        return {"available": False}
    if not manifest_path.is_file() or not prompt_path.is_file():
        raise ValueError("The global roster instruction and manifest must exist together")
    manifest = read_json(manifest_path)
    if not isinstance(manifest, dict) or manifest.get("schemaVersion") != 1:
        raise ValueError("Unsupported global roster manifest")
    expected_inputs = {
        "plan": PLAN, "bible": BIBLE,
        "pilot": ROOT / "public/cg/cg-manifest.json",
        "roster": ROOT / "src/game/cg/cgRosterArtwork.generated.json",
        "revisions": ROOT / "src/game/cg/cgRevisions.generated.json",
    }
    for name, source in expected_inputs.items():
        pin = manifest.get("inputs", {}).get(name, {})
        if pin.get("path") != source.relative_to(ROOT).as_posix() or not source.is_file():
            raise ValueError(f"Missing global roster input provenance: {name}")
        if pin.get("sha256") != hashlib.sha256(source.read_bytes()).hexdigest():
            raise ValueError(f"Stale global roster input: {name}")
    instruction = manifest.get("instruction", {})
    prompt_sha = hashlib.sha256(prompt_path.read_bytes()).hexdigest()
    if (instruction.get("path") != prompt_path.relative_to(ROOT).as_posix()
            or instruction.get("sha256") != prompt_sha):
        raise ValueError("The common roster instruction differs from its manifest")
    expected = {subject["characterId"]: subject for subject in plan["subjects"]}
    actual = manifest.get("subjects") or []
    if (len(actual) != len(expected)
            or {subject.get("characterId") for subject in actual} != set(expected)):
        raise ValueError("The global manifest must cover every roster subject once")
    slots = {slot["slot"] for slot in plan["slots"]}
    queue_counts = Counter()
    jobs_count = 0
    for subject in actual:
        jobs = subject.get("jobs") or []
        if len(jobs) != len(slots) or {job.get("slot") for job in jobs} != slots:
            raise ValueError(f"The global manifest lost a CG slot: {subject['characterId']}")
        for job in jobs:
            if job.get("production") != expected[subject["characterId"]]["matrix"][job["slot"]]:
                raise ValueError(f"Global queue changed actual production data: {subject['characterId']}/{job['slot']}")
            if not isinstance(job.get("queueStatus"), str) or not job["queueStatus"]:
                raise ValueError("Every global job needs an explicit queue disposition")
            queue_counts[job["queueStatus"]] += 1
            jobs_count += 1
    counts = manifest.get("summary", {})
    required_counts = {
        "subjects": len(expected), "universes": plan["summary"]["universes"],
        "cgSlots": jobs_count, "actualDeliveredCg": len(current),
        "documentedNotApplicable": plan["summary"]["matrixNotApplicableTotal"],
        "remainingSlots": plan["summary"]["totalRequestedSlotsNotProduced"],
        "automaticPerfectFidelityApprovals": 0,
        "imageCallsMadeByThisBuilder": 0, "imagesInstalledByThisBuilder": 0,
    }
    if any(counts.get(key) != value for key, value in required_counts.items()):
        raise ValueError("Global queue counters disagree with actual production")
    if counts.get("queueStatuses") != dict(queue_counts):
        raise ValueError("Global queue disposition counters do not match its jobs")
    if not manifest.get("rules", {}).get("queuedWorkIsNotDeliveredArtwork"):
        raise ValueError("Global queue must remain distinct from delivered artwork")
    return {
        "available": True,
        "manifestPath": manifest_path.relative_to(ROOT).as_posix(),
        "manifestSha256": hashlib.sha256(manifest_path.read_bytes()).hexdigest(),
        "masterPromptPath": prompt_path.relative_to(ROOT).as_posix(),
        "masterPromptSha256": prompt_sha,
        "summary": counts,
    }


def build(output: Path, no_images=False) -> dict:
    global LINK_BASE
    LINK_BASE = output.parent
    plan = read_json(PLAN)
    current_plan_entries = {entry["id"]: entry for entry in plan.get("entries", [])}
    bible_doc = read_json(BIBLE, {})
    bible = {row["characterId"]: row for row in bible_doc.get("subjects", bible_doc.get("characters", []))}
    manual = read_manual(output)
    versions, entry_by_id = collect_versions(plan)
    current = {key: values[-1] for key, values in versions.items() if values}
    visual_evidence = collect_visual_evidence(bible)
    production_notes = collect_production_notes()
    subjects = plan["subjects"]
    type_slot = {t: slot["slot"] for slot in plan["slots"] for t in slot["types"]}
    cg_by_subject_slot = {}
    for cg_id, entry in entry_by_id.items():
        char_id = entry.get("characterId") or entry.get("heroId")
        if char_id and entry.get("type") in type_slot:
            cg_by_subject_slot[(char_id, type_slot[entry["type"]])] = cg_id
    backgrounds = {entry["id"]: entry for entry in plan["backgrounds"]}
    roster_rows = {s["characterId"]: i + 4 for i, s in enumerate(subjects)}
    matrix_rows = {s["characterId"]: i * 15 + 4 for i, s in enumerate(subjects)}
    states = Counter()
    delivered_by_subject = Counter()
    na_by_subject = Counter()
    for subject in subjects:
        for slot_info in plan["slots"]:
            slot = slot_info["slot"]
            item = subject.get("matrix", {}).get(slot, {})
            cg_id = cg_by_subject_slot.get((subject["characterId"], slot)) or item.get("entryId")
            if cg_id in current:
                states["Livrée · revue 1:1 à compléter"] += 1
                delivered_by_subject[subject["characterId"]] += 1
            elif str(item.get("status", "")).startswith("not-applicable"):
                states["N/A documenté"] += 1
                na_by_subject[subject["characterId"]] += 1
            else:
                states["À produire"] += 1
    global_batch = collect_global_batch(plan, current)

    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = "Multiverse Breach — Fidélité visuelle et CG"
    wb.properties.subject = "Roster complet, cohérence des univers et suivi réel des artworks"
    wb.properties.creator = "Multiverse Breach"
    wb.properties.description = "Objectif 1:1 à vérifier par axes; absence de certification automatique."
    wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True)
    streams = []

    home = wb.create_sheet("Accueil")
    home.sheet_view.showGridLines = False
    home.sheet_properties.pageSetUpPr.fitToPage = True
    home.page_setup.orientation = "landscape"
    home.page_setup.paperSize = home.PAPERSIZE_A4
    home.page_setup.fitToWidth = 1
    home.page_setup.fitToHeight = 1
    home.print_area = "A1:F48" if global_batch["available"] else "A1:F38"
    home.column_dimensions["A"].width = 4
    home.column_dimensions["B"].width = 47
    home.column_dimensions["C"].width = 24
    home.column_dimensions["D"].width = 24
    home.column_dimensions["E"].width = 24
    home.column_dimensions["F"].width = 24
    home.merge_cells("B2:F3")
    home["B2"] = "MULTIVERSE BREACH\nFIDÉLITÉ VISUELLE & CG"
    home["B2"].fill = PatternFill("solid", fgColor=BLUE)
    home["B2"].font = Font(name="Aptos Display", size=23, bold=True, color=WHITE)
    home["B2"].alignment = Alignment(wrap_text=True, vertical="center")
    home.row_dimensions[2].height = 32
    home.row_dimensions[3].height = 32
    home.merge_cells("B5:F6")
    home["B5"] = "Objectif : même personnage, incarnation et univers cohérents, avec une pose et un décor distincts par CG. Les vignettes montrent les fichiers du projet. La fidélité se contrôle avec une source qualifiée : référence externe ou canon interne d'une création originale."
    home["B5"].font = Font(name="Aptos", size=12, color=INK)
    home["B5"].alignment = Alignment(wrap_text=True, vertical="center")
    home.row_dimensions[5].height = 40
    home.row_dimensions[6].height = 40
    metrics = [
        ("Personnages inclus", len(subjects)),
        ("Univers inclus", len({s['universe'] for s in subjects})),
        ("Emplacements CG (15 par personnage)", len(subjects) * 15),
        ("CG livrées (fichiers existants)", len(current)),
        ("CG restant à produire", states["À produire"]),
        ("N/A documentées", states["N/A documenté"]),
        ("Versions conservées, actuelles incluses", sum(map(len, versions.values()))),
        ("Fiches de cohérence issues de la bible", len(bible)),
        ("Personnages aux CG applicables complètes", sum(delivered_by_subject[s["characterId"]] > 0 and delivered_by_subject[s["characterId"]] + na_by_subject[s["characterId"]] == 15 for s in subjects)),
        ("Part des emplacements CG livrés", len(current) / (len(subjects) * 15)),
    ]
    for row, (label, number) in enumerate(metrics, 8):
        home.cell(row, 2, label)
        home.cell(row, 3, number)
        home.cell(row, 2).font = Font(name="Aptos", color=INK, size=12)
        home.cell(row, 3).font = Font(name="Aptos", bold=True, color=BLUE, size=15)
        home.cell(row, 3).number_format = "0.00%" if isinstance(number, float) else "#,##0"
        home.cell(row, 2).fill = home.cell(row, 3).fill = PatternFill("solid", fgColor=LIGHT if row % 2 == 0 else WHITE)
        home.row_dimensions[row].height = 25
    home["E14"] = "Vue filtrée du Roster"
    home["E14"].font = Font(name="Aptos", bold=True, color=BLUE, size=11)
    last_roster_row = len(subjects) + 3
    for row, label, formula in [
        (15, "Personnages visibles", f"=SUBTOTAL(103,'Roster'!B4:B{last_roster_row})"),
        (16, "CG livrées visibles", f"=SUBTOTAL(109,'Roster'!Z4:Z{last_roster_row})"),
        (17, "CG à produire visibles", f"=SUBTOTAL(109,'Roster'!AA4:AA{last_roster_row})"),
        (18, "N/A visibles", f"=SUBTOTAL(109,'Roster'!AB4:AB{last_roster_row})"),
    ]:
        home.cell(row, 5, label)
        home.cell(row, 6, formula)
        home.cell(row, 5).font = BODY_FONT
        home.cell(row, 6).font = Font(name="Aptos", bold=True, color=BLUE, size=12)
        home.cell(row, 6).number_format = "#,##0"
    home.merge_cells("B19:F20")
    home["B19"] = "Mode d'emploi : filtrer le Roster par univers ou priorité, ouvrir les références, vérifier les 8 axes dans Contrôles, puis suivre chaque CG dans Matrice CG. Le statut Livrée confirme un fichier enregistré; la revue 1:1 reste indépendante."
    home["B19"].alignment = Alignment(wrap_text=True, vertical="center")
    home["B19"].font = Font(name="Aptos", size=11, color=INK)
    home.row_dimensions[19].height = home.row_dimensions[20].height = 40
    for row, (sheet, label) in enumerate([
        ("Roster", "1912 fiches personnages et visuels réels"),
        ("Contrôles", "Revue manuelle de fidélité, axe par axe"),
        ("Matrice CG", "Toutes les CG, thèmes, poses, décors et états"),
        ("CG livrées", "Planches des images actuellement livrées"),
        ("Comparaisons", "Source indépendante, visuel projet et CG côte à côte"),
        ("Variantes", "Versions historiques conservées"),
        ("Variantes sprites", "Sprites actuels et anciennes variantes conservées"),
        ("Univers", "Cohérence par univers"),
        ("Sources", "Références et niveau de preuve"),
        ("Règles CG", "Thèmes imposés et limites des transformations"),
        ("Traçabilité", "Données sources, méthode et reconstruction"),
    ], 23):
        link_cell(home.cell(row, 2), f"#'{sheet}'!A1", label)
        home.row_dimensions[row].height = 23
    home.merge_cells("B36:F37")
    home["B36"] = "Les décisions de revue et commentaires manuels sont conservés au prochain rebuild. Les états réels de livraison et chemins d'assets sont recalculés depuis les registres du projet. Ne pas utiliser une vignette générée comme preuve de conformité canon."
    home["B36"].alignment = Alignment(wrap_text=True, vertical="center")
    home["B36"].font = Font(name="Aptos", italic=True, size=10, color=MUTED)
    home.row_dimensions[36].height = home.row_dimensions[37].height = 24
    if global_batch["available"]:
        home.merge_cells("B40:F40")
        home["B40"] = "UNE CONSIGNE COMMUNE · TOUT LE ROSTER"
        home["B40"].font = Font(name="Aptos", bold=True, size=12, color=BLUE)
        home.merge_cells("B41:F41")
        link_cell(home["B41"], global_batch["masterPromptPath"], "Ouvrir la consigne unique de fidélité et de production")
        home.merge_cells("B42:F42")
        link_cell(home["B42"], global_batch["manifestPath"], "Ouvrir les 1 912 personnages et leurs 28 680 travaux liés aux références")
        home.merge_cells("B44:F46")
        home["B44"] = "La consigne est commune ; chaque travail garde sa propre référence, sa pose et son décor. Le manifeste distingue qualification des sources, scènes à préciser, refus et livraisons. Une entrée de file n'est pas une CG produite ; aucun ordonnanceur ni lancement automatique n'est configuré."
        home["B44"].alignment = Alignment(wrap_text=True, vertical="center")
        home["B44"].font = Font(name="Aptos", size=11, color=INK)
        for row in [41, 42, 44, 45, 46]:
            home.row_dimensions[row].height = 24
    home.freeze_panes = "B8"

    roster = base_sheet(wb, "Roster", ROSTER_HEADERS, "Toutes les identités du roster. Vignette = fichier actuel du projet, jamais preuve canon par défaut. Les colonnes de contrôle renvoient à une revue manuelle.", "F4")
    controls = base_sheet(wb, "Contrôles", CONTROL_HEADERS, "8 axes indépendants. Conforme exige une comparaison visuelle avec la bonne incarnation. Aucune somme ni formule ne certifie automatiquement une fidélité 1:1.")
    matrix = base_sheet(wb, "Matrice CG", MATRIX_HEADERS, "15 emplacements par personnage. Livrée exige un fichier réel enregistré. Les thèmes What If sont identifiés; pose et décor doivent varier sans perdre l'identité source.", "F4")
    delivered = base_sheet(wb, "CG livrées", ["Image actuelle", "ID personnage", "Personnage", "Univers", "CG", "Type", "Révision", "Thème", "Pose", "Décor", "Revue visuelle", "Limites connues", "Image WebP", "PNG source", "Miniature", "Sources", "Aller matrice"], "Images réellement enregistrées et existantes. Proximité visuelle et fidélité 1:1 ne sont pas synonymes; les limites connues restent visibles.", "E4")
    variants_ws = base_sheet(wb, "Variantes", ["ID CG", "Personnage", "Univers", "CG", "Révision", "Version actuelle", "Date publication", "Image WebP", "PNG source", "Miniature", "Pose", "Thème", "ID décor", "État revue"], "Versions historiques et révisions enregistrées. Le classeur ne déplace ni ne remplace aucun ancien fichier.")
    universes_ws = base_sheet(wb, "Univers", ["Univers", "Personnages", "CG livrées", "CG à produire", "N/A documentées", "Œuvres / continuités", "Époques", "Lieux", "Architecture", "Palette", "Accessoires", "Limites de cohérence", "Exemples personnages", "Fiches roster"], "Cohérence de l'univers à appliquer avant génération. Les propositions de décor restent des propositions quand elles ne sont pas documentées par une source.")
    sources_ws = base_sheet(wb, "Sources", ["ID personnage", "Personnage", "Univers", "Référence", "Rôle", "Statut de preuve", "Fichier local existant", "Ouvrir", "Aller au roster"], "Sources textuelles, références visuelles externes et assets du projet restent distingués. Une URL seule ne prouve pas qu'une image a été contrôlée.")
    comparisons = base_sheet(wb, "Comparaisons", ["Référence source", "Visuel projet / planche", "CG actuelle", "Personnage", "ID personnage", "Incarnation", "URL source", "Fichier source local", "Statut de la source", "Consultation documentée", "Axes à comparer", "Limites / différences connues", "Reçu documentaire", "Contrôles", "Usage de la référence / CG"], "Comparer une source externe réelle ou une autorité canon interne explicitement documentée à la création actuelle. Une référence de costume croisé est limitée au créneau What If indiqué; elle ne prouve pas l'identité du personnage.", "F4")
    sprite_variants = base_sheet(wb, "Variantes sprites", ["Visuel projet", "ID personnage", "Personnage", "Univers", "Sélection actuelle du jeu", "Fichier sprite", "Rôle documentaire", "Statut de fidélité", "Roster"], "Personnages possédant plusieurs sprites locaux. Les anciennes images restent accessibles comme variantes; aucun de ces assets générés n'est une source canon indépendante.", "E4")

    sources_count = 0
    roster_images = 0
    cg_images = 0
    comparison_source_images = 0
    internal_source_images = 0
    comparison_project_images = 0
    sprite_variant_images = 0
    for subject in subjects:
        char_id = subject["characterId"]
        dossier = bible.get(char_id, {})
        identity = subject.get("identity", {})
        incarnation = dossier.get("incarnation") or {}
        if not isinstance(incarnation, dict):
            incarnation = {"label": incarnation}
        locks = dossier.get("visualLocks") or {}
        coherent = dossier.get("universeCoherence") or {}
        readiness = dossier.get("readiness") or {}
        row_num = roster_rows[char_id]
        cg_count = delivered_by_subject[char_id]
        na_count = na_by_subject[char_id]
        visual, visual_role = subject_visual(subject, dossier, current, entry_by_id)
        refs = dossier.get("sourceRefs") or identity.get("sourceRefs") or []
        if not isinstance(refs, list):
            refs = [refs]
        refs = list(refs)
        known_refs = {(ref.get("url") or ref.get("path")) if isinstance(ref, dict) else str(ref) for ref in refs}
        for evidence in visual_evidence.get(char_id, []):
            for key in ["url", "path"]:
                value = evidence.get(key)
                if value and value not in known_refs:
                    refs.append({key: value, "role": evidence.get("role") or "independent-visual-evidence", "status": "consultée selon reçu; portée et ressemblance à contrôler" if evidence.get("viewed") is True else evidence.get("status") or "image externe à qualifier"})
                    known_refs.add(value)
        reference_status = readiness.get("visualReference") or identity.get("evidenceStatus")
        if visual_evidence.get(char_id):
            internal_count = sum(e.get("sourceScope") == "original-project-canon" for e in visual_evidence[char_id])
            costume_count = sum(e.get("referencePurpose") == "costume-only" for e in visual_evidence[char_id])
            external_count = len(visual_evidence[char_id]) - internal_count - costume_count
            reference_status = f"{external_count} référence(s) externe(s), {internal_count} canon interne, {costume_count} tenue croisée seulement; comparaison 1:1 à compléter.\nStatut de la bible: {plain(reference_status)}"
        sprite_review = readiness.get("currentSprite") or {}
        sprite_status = sprite_review.get("status") if isinstance(sprite_review, dict) else None
        sprite_receipt = sprite_review.get("evidenceFile") if isinstance(sprite_review, dict) else None
        sprite_notes = [readiness.get("missing") or identity.get("missing")]
        visual_role_details = [visual_role]
        if sprite_status:
            visual_role_details.append("Sprite courant: " + plain(sprite_status))
            sprite_notes += [sprite_review.get("observation"), sprite_review.get("observations"), sprite_review.get("limitations")]
            if sprite_receipt:
                visual_role_details.append("Reçu de revue: " + plain(sprite_receipt))
        append_row(roster, [
            "" if not no_images else "Vignette omise (--no-images)", char_id, subject["characterName"], subject["universe"], subject.get("priority"),
            incarnation.get("label") or identity.get("incarnation") or "À déterminer depuis les sources",
            incarnation.get("status") or readiness.get("identity") or identity.get("evidenceStatus"),
            locks.get("silhouette") or dossier.get("identityLock") or identity.get("identityLock") or "À contrôler",
            locks.get("anatomy") or "À contrôler", locks.get("face") or "À contrôler",
            locks.get("costume") or locks.get("outfit") or "À contrôler", locks.get("palette") or "À contrôler",
            locks.get("equipment") or identity.get("equipment") or "À contrôler",
            locks.get("distinguishingMarks") or "À contrôler", locks.get("forbiddenMixes") or "Ne pas mélanger plusieurs incarnations sans thème What If explicite.",
            coherent.get("sourceWork") or subject["universe"], coherent.get("period") or "À préciser",
            coherent.get("locations"), coherent.get("architecture"), coherent.get("palette"), coherent.get("props"), coherent.get("boundaries"),
            reference_status, readiness.get("sourceReview") or "À contrôler",
            "Non certifiée — revue par axes requise", cg_count, 15 - cg_count - na_count, na_count,
            visual, ref_lines(refs), sprite_notes, "Ouvrir les 15 CG", "Vérifier les axes", visual_role_details,
        ])
        link_cell(roster.cell(row_num, 29), visual)
        link_cell(roster.cell(row_num, 32), f"#'Matrice CG'!A{matrix_rows[char_id]}")
        link_cell(roster.cell(row_num, 33), f"#'Contrôles'!A{row_num}")
        if isinstance(sprite_receipt, str) and local_path(sprite_receipt):
            link_cell(roster.cell(row_num, 34), sprite_receipt)
        if not no_images:
            roster_images += int(embed(roster, f"A{row_num}", local_path(visual), streams, (110, 110)))

        preserved = manual["Contrôles"].get((char_id,), {})
        control_values = [char_id, subject["characterName"], subject["universe"]]
        control_values += [preserved.get(axis) or "À contrôler" for axis in AXES]
        control_values += [None, "Non certifiée", preserved.get("Décision manuelle") or "Non revue", preserved.get("Commentaire QA"), preserved.get("Nom du relecteur"), preserved.get("Date de revue"), "Ouvrir"]
        control_values += [preserved.get("Cible d'incarnation manuelle"), preserved.get("Référence visuelle validée (URL / fichier)")]
        append_row(controls, control_values)
        controls.cell(row_num, 12, f'=COUNTIF(D{row_num}:K{row_num},"Conforme")')
        # Even 8/8 axes does not automatically convert into a certification.
        controls.cell(row_num, 13, f'=IF(L{row_num}=8,"8 axes conformes · décision manuelle requise","À contrôler · non certifiée")')
        link_cell(controls.cell(row_num, 18), f"#'Roster'!A{row_num}")

        source_refs = []
        for ref in refs:
            source_refs.append(ref if isinstance(ref, dict) else {"url" if str(ref).startswith("http") else "path": ref, "role": "source héritée du plan", "status": "preuve à qualifier"})
        for candidate in identity.get("localCandidates", []):
            source_refs.append({"path": candidate.get("path"), "role": candidate.get("role") or "asset visuel projet", "status": "image du projet — pas une référence canon indépendante"})
        seen_refs = set()
        for source in source_refs:
            ref = source.get("url") or source.get("path") or ""
            if not ref or ref in seen_refs:
                continue
            seen_refs.add(ref)
            append_row(sources_ws, [char_id, subject["characterName"], subject["universe"], ref, source.get("role"), source.get("status") or "À qualifier", bool(local_path(ref)), "Ouvrir", "Roster"])
            sr = sources_ws._current_row
            link_cell(sources_ws.cell(sr, 8), ref)
            link_cell(sources_ws.cell(sr, 9), f"#'Roster'!A{row_num}")
            sources_count += 1

        for slot_info in plan["slots"]:
            slot = slot_info["slot"]
            item = subject.get("matrix", {}).get(slot, {})
            type_name = item.get("type") or item.get("selectedType") or " / ".join(item.get("typeOptions") or slot_info["types"])
            cg_id = cg_by_subject_slot.get((char_id, slot)) or item.get("entryId")
            actual = current.get(cg_id, {})
            actual_entry = entry_by_id.get(cg_id, {})
            if actual:
                type_name = actual_entry.get("type") or type_name
            theme = theme_for(dossier, slot, type_name)
            bg_id = actual.get("backgroundId") or item.get("backgroundId")
            bg = backgrounds.get(bg_id, {})
            planned_background = theme.get("backgroundBrief") or theme.get("background") or bg.get("description")
            if not planned_background:
                env = backgrounds.get(item.get("environmentSourceId"), {})
                planned_background = "Décor proposé, à sourcer: " + plain(env.get("sourceSetting") or coherent.get("locations") or subject["universe"])
            background = actual.get("backgroundDescription") or actual.get("background") or planned_background
            real_state = "Livrée · revue 1:1 à compléter" if actual else "N/A documenté" if str(item.get("status", "")).startswith("not-applicable") else "À produire"
            production_note = production_notes.get((char_id, slot), {})
            if production_note.get("blocked"):
                real_state = "Ancienne image · correction refusée par le service" if actual else "À produire · refus du service de génération"
            elif (actual.get("visualReview") or {}).get("status") == "needs-source-correction":
                real_state = "Ancienne image · correction source nécessaire"
            if real_state.startswith("N/A"):
                theme_brief = "N/A justifié: " + plain(item.get("reason") or item.get("source"))
            else:
                # The common plan expands terse receipt labels from their
                # documented prompt summary. Reuse that current description;
                # historical variant rows still retain their own metadata.
                current_plan_theme = current_plan_entries.get(cg_id, {}).get("themeBrief") if actual else None
                theme_brief = current_plan_theme or actual.get("themeBrief") or theme.get("themeBrief") or theme.get("brief") or item.get("themeBrief")
            variant_count = max(0, len(versions.get(cg_id, [])) - 1)
            preserved_cg = manual["Matrice CG"].get((char_id, slot), {})
            family = actual_entry.get("family")
            nature = family or ("What If / transformation imposée" if int(slot[-2:]) >= 12 else "Fan Art / thème imposé" if int(slot[-2:]) >= 9 else "Action Nexus · adaptation documentée" if int(slot[-2:]) == 5 else "Canon / scène cohérente")
            append_row(matrix, [
                char_id, subject["characterName"], subject["universe"], subject.get("priority"), slot, type_name,
                nature,
                theme_brief, actual.get("poseBrief") or actual.get("poseDescription") or theme.get("poseBrief") or item.get("poseBrief"), background, bg_id,
                real_state, actual.get("themeStatus") or theme.get("status") or item.get("themeStatus") or "À qualifier",
                incarnation.get("label") or identity.get("incarnation") or "À déterminer", actual.get("themeSource") or theme.get("source") or item.get("themeSource") or item.get("source"),
                ref_lines(actual.get("sourceRefs") or refs), actual.get("imagePath"), actual.get("sourcePath"), variant_count,
                compact_review(actual.get("visualReview")) if actual else "À contrôler après production",
                [(actual.get("visualReview") or {}).get("limitations", []) if isinstance(actual.get("visualReview"), dict) else "", production_note.get("text") or ""],
                preserved_cg.get("Revue manuelle") or "Non revue", preserved_cg.get("Commentaire QA"), "Roster",
            ])
            mr = matrix._current_row
            link_cell(matrix.cell(mr, 17), actual.get("imagePath"))
            link_cell(matrix.cell(mr, 18), actual.get("sourcePath"))
            current_review = actual.get("visualReview") or {}
            independent_receipt = (current_review.get("independentReviewFile") or current_review.get("independentReview")) if isinstance(current_review, dict) else None
            if isinstance(independent_receipt, str) and (local_path(independent_receipt) or Path()).is_file():
                link_cell(matrix.cell(mr, 20), independent_receipt)
            link_cell(matrix.cell(mr, 24), f"#'Roster'!A{row_num}")
            if actual:
                review = actual.get("visualReview") or {}
                append_row(delivered, ["", char_id, subject["characterName"], subject["universe"], slot, type_name,
                    actual.get("revisionId"), theme_brief, actual.get("poseBrief"), background, compact_review(review),
                    review.get("limitations") or review.get("notes") if isinstance(review, dict) else "",
                    actual.get("imagePath"), actual.get("sourcePath"), actual.get("thumbnailPath"), ref_lines(actual.get("sourceRefs") or refs), "Matrice"])
                dr = delivered._current_row
                if isinstance(independent_receipt, str) and (local_path(independent_receipt) or Path()).is_file():
                    link_cell(delivered.cell(dr, 11), independent_receipt)
                for col, key in [(13, "imagePath"), (14, "sourcePath"), (15, "thumbnailPath")]:
                    link_cell(delivered.cell(dr, col), actual.get(key))
                link_cell(delivered.cell(dr, 17), f"#'Matrice CG'!A{mr}")
                if not no_images:
                    cg_images += int(embed(delivered, f"A{dr}", local_path(actual.get("thumbnailPath") or actual.get("imagePath")), streams, (90, 120)))
                for version in versions[cg_id]:
                    append_row(variants_ws, [cg_id, subject["characterName"], subject["universe"], slot, version.get("revisionId"), version is actual, version.get("publishedAt"), version.get("imagePath"), version.get("sourcePath"), version.get("thumbnailPath"), version.get("poseBrief"), theme_brief if version is actual else version.get("themeBrief"), version.get("backgroundId"), compact_review(version.get("visualReview"))])
                    vr = variants_ws._current_row
                    for col, key in [(8, "imagePath"), (9, "sourcePath"), (10, "thumbnailPath")]:
                        link_cell(variants_ws.cell(vr, col), version.get(key))

    universe_groups = defaultdict(list)
    for subject in subjects:
        universe_groups[subject["universe"]].append(subject)
    for universe, members in sorted(universe_groups.items()):
        def group_coherence(field):
            values = []
            for s in members:
                value = plain((bible.get(s["characterId"], {}).get("universeCoherence") or {}).get(field))
                if value and value not in values:
                    values.append(value)
            return "\n\n".join(values)
        count = sum(delivered_by_subject[s["characterId"]] for s in members)
        na_count = sum(na_by_subject[s["characterId"]] for s in members)
        append_row(universes_ws, [universe, len(members), count, len(members) * 15 - count - na_count, na_count,
            group_coherence("sourceWork"), group_coherence("period"), group_coherence("locations"), group_coherence("architecture"), group_coherence("palette"), group_coherence("props"), group_coherence("boundaries"),
            "\n".join(s["characterName"] for s in members[:12]), "Ouvrir le premier personnage"])
        link_cell(universes_ws.cell(universes_ws._current_row, 14), f"#'Roster'!A{roster_rows[members[0]['characterId']]}")

    for subject in subjects:
        char_id = subject["characterId"]
        # All produced identities and all priority corrections receive a direct
        # comparison row, even when the external bitmap is still missing.
        if not delivered_by_subject[char_id] and subject.get("priority") != "P0" and char_id not in visual_evidence:
            continue
        dossier = bible.get(char_id, {})
        incarnation = dossier.get("incarnation") or {}
        incarnation_label = incarnation.get("label") if isinstance(incarnation, dict) else incarnation
        original_project = isinstance(incarnation, dict) and incarnation.get("status") == "original-project-canon"
        local_candidates = [v for v in dossier.get("localVisuals", []) if isinstance(v, dict) and v.get("selectedInRuntime")]
        local_candidates += subject.get("identity", {}).get("localCandidates", [])
        project_path = next((c.get("path") for c in local_candidates if local_path(c.get("path"))), None)
        references = visual_evidence.get(char_id) or [{}]
        if references and all(e.get("referencePurpose") in {"costume-only", "equipment-only", "environment-only", "body-costume-support"} for e in references):
            references = [{}, *references]
        for evidence in references:
            target_slot = evidence.get("slot") or "CG01"
            target_slot = str(target_slot) if str(target_slot).startswith("CG") else f"CG{int(target_slot):02}"
            current_cg = current.get(cg_by_subject_slot.get((char_id, target_slot)), {})
            evidence_path = local_path(evidence.get("path"))
            evidence_status = evidence.get("status") or evidence.get("provenance") or evidence.get("role") or "Source externe téléchargée; portée à contrôler" if evidence_path else "Aucun bitmap source indépendant disponible"
            if evidence.get("sourceScope") == "original-project-canon":
                authority_role = evidence.get("authorityRole") or evidence.get("role")
                authority_kind = "portrait original" if authority_role == "canonical-original-project-portrait" else "décor original" if authority_role == "original-project-environment-reference" else "pack original"
                authority_owner = "de l'univers" if authority_role == "original-project-environment-reference" else "du personnage"
                evidence_status = f"Autorité canon interne documentée : {authority_kind} {authority_owner}; pas une source externe"
                if evidence.get("reusedAuthorityBitmap"):
                    evidence_status += ". Réutilisation du même bitmap source pour ce créneau uniquement; aucune nouvelle source ni preuve d'identité indépendante."
                if evidence.get("inspectionPath"):
                    evidence_status += ". Vignette : détail agrandi du même pack source."
            elif evidence.get("inspectionPath"):
                evidence_status += ". Vignette : détail agrandi du même bitmap source."
            viewed = evidence.get("viewed")
            consultation = "Consultée selon reçu documentaire" if viewed is True else "Non consultée selon reçu" if viewed is False else "Consultation non explicitement documentée"
            if viewed is None and "historically-viewed" in str(evidence.get("status", "")):
                consultation = "Consultée historiquement selon reçu; comparaison à reprendre"
            if not evidence_path:
                consultation = "Référence à récupérer / qualifier"
                if original_project:
                    evidence_status = "Création originale du projet : autorité canon interne à contrôler"
                    consultation = "Comparer à la référence interne retenue; aucune source externe requise par principe"
            axes = (dossier.get("visualLocks") or {})
            purpose = evidence.get("referencePurpose")
            usage = "identité / incarnation"
            if purpose == "costume-only":
                axes = "Tenue croisée uniquement, " + target_slot + ". Ne pas utiliser cette image pour valider visage, silhouette ou anatomie du personnage."
                usage = "costume croisé What If seulement"
            elif purpose == "equipment-only":
                axes = "Équipement uniquement, " + target_slot + ". Comparer forme, matériaux et éléments distinctifs de l'objet; cette référence ne valide pas le visage ou la tenue du personnage."
                usage = "équipement uniquement"
            elif purpose == "environment-only":
                axes = "Décor / univers uniquement, " + target_slot + ". Comparer architecture, époque, palette et accessoires du lieu; les autres personnages éventuellement visibles ne servent pas de preuve d'identité."
                usage = "décor / univers uniquement"
            elif purpose == "body-costume-support":
                axes = "Support de silhouette et tenue du corps uniquement. La figurine dérivée complète les cadrages du film sans certifier le visage, les articulations ou tous les accessoires du modèle original. Ce n'est pas une tenue croisée CG13."
                usage = "corps / tenue — support dérivé, pas une CG13"
            limits = evidence.get("limitations") or (current_cg.get("visualReview") or {}).get("limitations") or (dossier.get("readiness") or {}).get("missing")
            source_sprite_review = (dossier.get("readiness") or {}).get("currentSprite") or {}
            if isinstance(source_sprite_review, dict) and source_sprite_review.get("status"):
                limits = [*(limits if isinstance(limits, list) else [limits] if limits else []), "Sprite courant: " + plain(source_sprite_review["status"]), source_sprite_review.get("observation"), source_sprite_review.get("observations")]
            if evidence.get("inspectionPath"):
                limits = ["Vignette source cadrée et agrandie depuis le même fichier source; le lien ouvre l'original. Ce détail n'ajoute aucune preuve ou source indépendante.", *(limits if isinstance(limits, list) else [limits] if limits else [])]
            if evidence.get("reusedAuthorityBitmap"):
                limits = ["Même bitmap canon interne que la référence d'identité, réutilisé uniquement pour " + target_slot + ". Aucun autre visage, corps ou costume n'est validé par cet usage.", evidence.get("referenceScope"), *(limits if isinstance(limits, list) else [limits] if limits else [])]
            append_row(comparisons, ["" if evidence_path else "Canon original interne" if original_project else "Source manquante", "", "" if current_cg else target_slot + " à produire", subject["characterName"], char_id,
                incarnation_label or subject.get("identity", {}).get("incarnation") or "À préciser", evidence.get("url") or evidence.get("sourceUrl"), evidence.get("path"), evidence_status, consultation,
                axes or AXES, limits, evidence.get("receipt"), "Vérifier les axes", target_slot + " — " + usage])
            cr = comparisons._current_row
            link_cell(comparisons.cell(cr, 7), evidence.get("url") or evidence.get("sourceUrl"))
            link_cell(comparisons.cell(cr, 8), evidence.get("path"))
            link_cell(comparisons.cell(cr, 13), evidence.get("receipt"))
            link_cell(comparisons.cell(cr, 14), f"#'Contrôles'!A{roster_rows[char_id]}")
            if not no_images:
                embedded = int(embed(comparisons, f"A{cr}", local_path(evidence.get("inspectionPath")) or evidence_path, streams, (145, 180)))
                if evidence.get("sourceScope") == "original-project-canon":
                    internal_source_images += embedded
                else:
                    comparison_source_images += embedded
                comparison_project_images += int(embed(comparisons, f"B{cr}", local_path(project_path), streams, (145, 180)))
                comparison_project_images += int(embed(comparisons, f"C{cr}", local_path(current_cg.get("thumbnailPath") or current_cg.get("imagePath")), streams, (135, 180)))

    for subject in subjects:
        char_id = subject["characterId"]
        seen = set()
        actual_variants = []
        for visual in bible.get(char_id, {}).get("localVisuals", []):
            if not isinstance(visual, dict):
                continue
            if visual.get("role") == "canonical-original-project-sprite-pack":
                continue
            raw = visual.get("path", "")
            if not raw.startswith("/sprites/") or raw in seen or not local_path(raw):
                continue
            seen.add(raw)
            actual_variants.append(visual)
        if len(actual_variants) < 2:
            continue
        for visual in actual_variants:
            append_row(sprite_variants, ["", char_id, subject["characterName"], subject["universe"], bool(visual.get("selectedInRuntime")), visual["path"], visual.get("role") or "Sprite généré du projet", visual.get("status") or "Fidélité 1:1 non certifiée", "Roster"])
            vr = sprite_variants._current_row
            link_cell(sprite_variants.cell(vr, 6), visual["path"])
            link_cell(sprite_variants.cell(vr, 9), f"#'Roster'!A{roster_rows[char_id]}")
            if not no_images:
                sprite_variant_images += int(embed(sprite_variants, f"A{vr}", local_path(visual["path"]), streams, (110, 110)))

    finish_table(roster, "T_Roster", {1: 17, 2: 26, 3: 25, 4: 26, 5: 15, 25: 28, 26: 12, 27: 12, 28: 12, 29: 46, 30: 65, 31: 48, 32: 19, 33: 19, 34: 42}, 87)
    finish_table(controls, "T_Controles", {1: 28, 2: 26, 3: 28, 12: 14, 13: 43, 15: 60, 18: 18, 19: 60, 20: 80}, 46)
    finish_table(matrix, "T_MatriceCG", {1: 27, 2: 25, 3: 25, 4: 14, 5: 10, 6: 24, 8: 65, 9: 55, 10: 72, 11: 45, 12: 32, 13: 42, 14: 45, 15: 53, 16: 65, 17: 50, 18: 50, 19: 16, 20: 43, 21: 55, 23: 55, 24: 15}, 66)
    finish_table(delivered, "T_CGLivrees", {1: 15, 2: 27, 3: 25, 4: 25, 5: 10, 6: 25, 8: 60, 9: 55, 10: 65, 11: 50, 12: 60, 13: 55, 14: 55, 15: 55, 16: 65, 17: 15}, 93)
    finish_table(variants_ws, "T_Variantes", {1: 68, 2: 25, 3: 25, 4: 10, 5: 40, 6: 15, 7: 22, 8: 55, 9: 55, 10: 55, 11: 55, 12: 60, 13: 50, 14: 48}, 60)
    finish_table(universes_ws, "T_Univers", {1: 34, 2: 15, 3: 15, 4: 17, 5: 17, 6: 55, 7: 45, 8: 65, 9: 65, 10: 50, 11: 55, 12: 65, 13: 45, 14: 30}, 70)
    finish_table(sources_ws, "T_Sources", {1: 28, 2: 26, 3: 28, 4: 85, 5: 44, 6: 60, 7: 22, 8: 16, 9: 16}, 38)
    finish_table(comparisons, "T_Comparaisons", {1: 22, 2: 22, 3: 21, 4: 27, 5: 28, 6: 45, 7: 70, 8: 70, 9: 56, 10: 46, 11: 72, 12: 75, 13: 65, 14: 24, 15: 48}, 138)
    finish_table(sprite_variants, "T_VariantesSprites", {1: 17, 2: 28, 3: 28, 4: 28, 5: 22, 6: 85, 7: 48, 8: 65, 9: 18}, 87)
    # Link styling must survive the general table typography pass.
    for ws in [roster, controls, matrix, delivered, variants_ws, universes_ws, sources_ws, comparisons, sprite_variants]:
        for row in ws.iter_rows(min_row=4):
            for cell in row:
                if cell.hyperlink:
                    cell.font = LINK_FONT

    rules_ws = base_sheet(wb, "Règles CG", ["CG", "Types", "Thème imposé / objectif", "Pose distincte", "Décor", "Source / application", "Limites"], "Le pilote définit la grammaire CG. Son application à un nouveau personnage demande une adaptation cohérente et explicitement identifiée.", "C4")
    for template in plan["extensionTemplates"]:
        append_row(rules_ws, [template["id"], template.get("typeOptions"), template.get("themeBrief"), template.get("poseBrief"), template.get("backgroundOptions"), template.get("themeSource"), template.get("resolutionRule")])
    finish_table(rules_ws, "T_ReglesCG", {1: 12, 2: 30, 3: 75, 4: 70, 5: 75, 6: 55, 7: 80}, 90)

    evidence_ws = base_sheet(wb, "Traçabilité", ["Élément", "Valeur", "Règle / interprétation"], "Données et limites de cette reconstruction. Le classeur sert à contrôler la fidélité; il ne constitue pas un certificat canon.", "B4")
    timestamp = datetime.now(timezone.utc).isoformat(timespec="seconds")
    trace_rows = [
        ("Généré le (UTC)", timestamp, "Reconstruction locale sans appel à un service tiers."),
        ("Roster source", str(PLAN.relative_to(ROOT)), "Tous les subjects du plan sont inclus; aucun personnage ajouté par approximation."),
        ("Bible d'incarnations", str(BIBLE.relative_to(ROOT)) if BIBLE.is_file() else "Non encore disponible lors de cette reconstruction", "Les propositions et sources héritées gardent leur statut documentaire."),
        ("Registre des révisions", "src/game/cg/cgRevisions.generated.json", "Chaque livraison doit avoir un fichier réel enregistré."),
        ("Originaux du pilote", "public/cg/cg-manifest.json", "Les 83 versions historiques sont prises dans le manifeste immuable, même quand le plan pointe désormais sur des révisions courantes."),
        ("Nouveaux artworks roster", "src/game/cg/cgRosterArtwork.generated.json", "Seuls les assets réellement présents apparaissent comme livrés."),
        ("Script de reconstruction", "scripts/buildCgFidelityWorkbook.py", "python scripts/buildCgFidelityWorkbook.py"),
        ("Dépendances", "openpyxl >= 3.1 ; Pillow", "Le script ne télécharge aucune image."),
        ("Anciennes variantes", sum(map(len, versions.values())) - len(current), "Versions enregistrées non actuelles; aucun fichier historique n'est modifié."),
        ("Images actuelles avec correction source requise", sum((v.get("visualReview") or {}).get("status") == "needs-source-correction" for v in current.values()), "Écarts identifiés dans la revue enregistrée. Une image physique existante n'est pas présentée comme une correction achevée."),
        ("Vignettes roster", roster_images, "Miniatures de fichiers locaux réels, planche entière ou CG actuelle. Aucune source canon inventée."),
        ("Vignettes CG livrées", cg_images, "Images réellement livrées dans le projet."),
        ("Comparaisons : sources indépendantes", comparison_source_images, "Bitmaps de sources externes réelles incorporés dans la colonne réservée; aucune image du projet n'y est substituée."),
        ("Comparaisons : autorités canon internes", internal_source_images, "Uniquement les références de création originale avec un reçu canon interne explicite, distinguées des sources externes."),
        ("Variantes de sprites", sprite_variants.max_row - 3, "Sprites sélectionnés et anciennes variantes réellement présentes pour les personnages ayant plusieurs fichiers locaux."),
        ("N/A", states["N/A documenté"], "Uniquement les exceptions documentées avec justification dans le plan. Les autres cas restent à produire/qualifier."),
        ("Fidélité 1:1", "Objectif de contrôle, aucune certification automatique", "Les 8 axes, les sources indépendantes et une décision humaine doivent être examinés."),
        ("Créations originales", "Référence canon interne du projet", "Une création originale se contrôle contre sa référence interne documentée. Un asset du projet n'est admis comme source qu'avec un reçu d'autorité canon interne explicite; il reste distingué d'une source externe."),
        ("Familles CG", "Canon : 01–04, 06–08 ; Nexus : 05 ; Fan Art : 09–11 ; What If : 12–15", "Les transformations restent identifiées comme alternatives; elles ne modifient pas la référence canon CG01. La famille enregistrée d'une image livrée prime sur ce modèle."),
        ("Pose / décor", "Brief distinct par emplacement", "Un brief différent ne prouve pas que deux images produites auront des poses réellement différentes; il faut les comparer."),
        ("États manuels préservés", sum(len(v) for v in manual.values()), "Contrôles (8 axes, décision, commentaire, relecteur, date) et Matrice CG (revue/commentaire) sont conservés au rebuild."),
        ("Choix du relecteur", "Colonnes Cible d'incarnation manuelle et Référence visuelle validée", "Ces choix sont préservés dans Contrôles; ils ne réécrivent pas automatiquement les données canon ou les registres du jeu."),
        ("Excel / calculs", "Recalcul à l'ouverture", "Le bilan global est figé au rebuild. La vue filtrée du Roster utilise SUBTOTAL et suit les filtres Excel; les formules des axes se recalculent dans Excel."),
        ("Liens locaux", "Relatifs à l'emplacement du classeur", "Les vignettes PNG incorporées restent lisibles dans le fichier Excel seul. Les liens aux fichiers haute définition suivent l'arborescence locale du projet."),
    ]
    if global_batch["available"]:
        trace_rows += [
            ("Consigne unique du roster", global_batch["masterPromptPath"], "Même instruction pour tous les sujets, avec une liaison propre à l'incarnation et au créneau ; plusieurs appels réels peuvent être nécessaires."),
            ("Manifeste de travail global", global_batch["manifestPath"], "Couverture et empreintes vérifiées contre le plan, la bible et les registres. La file ne constitue pas une livraison."),
            ("Dispositions de la file globale", global_batch["summary"]["queueStatuses"], "Qualification / préparation des travaux uniquement ; les états réels et fichiers restent dans Matrice CG et CG livrées."),
            ("Appels images du builder de manifeste", global_batch["summary"]["imageCallsMadeByThisBuilder"], "Le générateur de manifeste n'appelle aucun service d'image, n'installe aucune CG et n'ajoute aucune certification."),
        ]
    for trace in trace_rows:
        append_row(evidence_ws, trace)
    finish_table(evidence_ws, "T_Tracabilite", {1: 36, 2: 90, 3: 100}, 45)
    if global_batch["available"]:
        for row in evidence_ws.iter_rows(min_row=4):
            if row[0].value in {"Consigne unique du roster", "Manifeste de travail global"}:
                link_cell(row[1], row[1].value)

    lists = wb.create_sheet("Listes")
    lists.append(["Axes", "Décision", "Revue CG"])
    axis_values = ["À contrôler", "Conforme", "À corriger", "N/A justifié"]
    decision_values = ["Non revue", "À corriger", "Prête pour revue finale", "Validée manuellement avec sources"]
    cg_values = ["Non revue", "À corriger", "Pose à différencier", "Décor à différencier", "Thème à corriger", "Prête pour revue finale", "Validée manuellement avec sources"]
    for i in range(max(map(len, [axis_values, decision_values, cg_values]))):
        lists.append([values[i] if i < len(values) else None for values in [axis_values, decision_values, cg_values]])
    lists.sheet_state = "hidden"
    for list_col, name, values in [(1, "FidelityAxisStates", axis_values), (2, "FidelityDecisions", decision_values), (3, "CgReviewStates", cg_values)]:
        letter = get_column_letter(list_col)
        wb.defined_names.add(DefinedName(name, attr_text=f"'Listes'!${letter}$2:${letter}${len(values) + 1}"))
    for col in range(4, 12):
        add_list_validation(controls, col, 1, len(axis_values) + 1)
    add_list_validation(controls, 14, 2, len(decision_values) + 1)
    add_list_validation(matrix, 22, 3, len(cg_values) + 1)
    for state, color in [("Conforme", GREEN), ("À corriger", RED), ("À contrôler", AMBER)]:
        controls.conditional_formatting.add(f"D4:K{controls.max_row}", CellIsRule(operator="equal", formula=[f'"{state}"'], fill=PatternFill("solid", fgColor=color)))
    for fragment, color in [("Livrée", GREEN), ("À produire", AMBER), ("N/A", LIGHT), ("Ancienne image", RED)]:
        matrix.conditional_formatting.add(f"L4:L{matrix.max_row}", FormulaRule(formula=[f'LEFT(L4,{len(fragment)})="{fragment}"'], fill=PatternFill("solid", fgColor=color)))
    roster.conditional_formatting.add(f"Z4:Z{roster.max_row}", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor=GREEN)))
    roster.conditional_formatting.add(f"AA4:AA{roster.max_row}", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor=AMBER)))
    home["E8"] = "Production CG"
    home["F8"] = "Emplacements"
    home["E9"] = "Livrées"
    home["F9"] = len(current)
    home["E10"] = "À produire"
    home["F10"] = states["À produire"]
    home["E11"] = "N/A documentées"
    home["F11"] = states["N/A documenté"]
    chart = BarChart()
    chart.title = "État réel des 28 680 emplacements"
    chart.y_axis.title = "Emplacements"
    chart.x_axis.title = "État"
    chart.add_data(Reference(home, min_col=6, min_row=8, max_row=11), titles_from_data=True)
    chart.set_categories(Reference(home, min_col=5, min_row=9, max_row=11))
    chart.height = 7.2
    chart.width = 16
    chart.legend = None
    home.add_chart(chart, "H8")
    wb.active = 0
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(output.stem + ".building.xlsx")
    wb.save(temporary)
    temporary.replace(output)
    summary = {
        "schemaVersion": 1, "generatedAt": timestamp,
        "workbook": str(output.relative_to(ROOT)) if output.is_relative_to(ROOT) else str(output),
        "subjects": len(subjects), "universes": len(universe_groups), "matrixRows": len(subjects) * 15,
        "deliveredCg": len(current), "unproducedSlots": states["À produire"], "documentedNotApplicable": states["N/A documenté"],
        "registeredVersions": sum(map(len, versions.values())), "preservedVariants": sum(map(len, versions.values())) - len(current),
        "bibleSubjects": len(bible), "rosterThumbnails": roster_images, "deliveredCgThumbnails": cg_images,
        "comparisonRows": comparisons.max_row - 3, "independentSourceThumbnails": comparison_source_images, "comparisonProjectThumbnails": comparison_project_images,
        "internalCanonSourceThumbnails": internal_source_images,
        "spriteVariantRows": sprite_variants.max_row - 3, "spriteVariantThumbnails": sprite_variant_images,
        "sourceRows": sources_count, "sheets": wb.sheetnames, "automaticFidelityCertifications": 0,
        "knownCurrentSourceCorrectionsRequired": sum((v.get("visualReview") or {}).get("status") == "needs-source-correction" for v in current.values()),
        "serviceRefusedSlots": sum(bool(v.get("blocked")) for v in production_notes.values()),
        "subjectsWithExternalIdentityReferences": sum(any(e.get("sourceScope") == "external-reference" and e.get("referencePurpose") not in {"costume-only", "equipment-only", "environment-only", "body-costume-support"} for e in evidence) for evidence in visual_evidence.values()),
        "subjectsWithInternalCanonAuthorities": sum(any(e.get("sourceScope") == "original-project-canon" and not e.get("reusedAuthorityBitmap") for e in evidence) for evidence in visual_evidence.values()),
        "scopedInternalBitmapReuses": sum(bool(e.get("reusedAuthorityBitmap")) for evidence in visual_evidence.values() for e in evidence),
        "subjectsWithCostumeOnlyReferences": sum(any(e.get("referencePurpose") == "costume-only" for e in evidence) for evidence in visual_evidence.values()),
        "workbookBytes": output.stat().st_size,
        "workbookSha256": hashlib.sha256(output.read_bytes()).hexdigest(),
        "manualRowsPreserved": {k: len(v) for k, v in manual.items()},
        "globalBatch": global_batch,
    }
    output.with_suffix(".summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return summary


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--no-images", action="store_true", help="Diagnostic only; the normal workbook includes real thumbnails.")
    args = parser.parse_args()
    print(json.dumps(build(args.output.resolve(), args.no_images), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
