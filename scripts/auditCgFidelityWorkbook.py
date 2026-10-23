#!/usr/bin/env python3
"""Independently verify the delivered Excel against the actual runtime media."""
from __future__ import annotations

import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
from urllib.parse import unquote
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / "docs/cg/Multiverse_Breach_Fidelite_CG.xlsx"
SLOTS = [
    ["characterSolo"], ["weaponSolo"], ["decorSolo"], ["coherentScene"],
    ["actionScene"], ["introPose"], ["victoryPose"], ["defeatPose"],
    ["beachFamily"], ["maidService"], ["goofy"], ["alignmentSwap"],
    ["genderSwap", "iconicOutfitSwap"], ["zombieVersion"],
    ["firstStep", "futureExperienced"],
]


def read_json(name):
    return json.loads((ROOT / name).read_text(encoding="utf-8"))


def rows(sheet):
    header = next(sheet.iter_rows(min_row=3, max_row=3, values_only=True))
    return (dict(zip(header, row)) for row in sheet.iter_rows(min_row=4, values_only=True))


def audit_global_batch(book, plan):
    summary = json.loads(book.with_suffix(".summary.json").read_text(encoding="utf-8"))
    batch = summary.get("globalBatch") or {}
    assert batch.get("available") is True, "The common instruction and complete roster manifest are missing"
    manifest_path = ROOT / batch["manifestPath"]
    prompt_path = ROOT / batch["masterPromptPath"]
    assert hashlib.sha256(manifest_path.read_bytes()).hexdigest() == batch["manifestSha256"]
    assert hashlib.sha256(prompt_path.read_bytes()).hexdigest() == batch["masterPromptSha256"]
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert manifest["instruction"]["path"] == batch["masterPromptPath"]
    assert manifest["instruction"]["sha256"] == batch["masterPromptSha256"]
    expected_inputs = {
        "plan": "docs/cg/roster-production-2026-10-18.json",
        "bible": "docs/cg/roster-fidelity-bible-2026-10-18.json",
        "pilot": "public/cg/cg-manifest.json",
        "roster": "src/game/cg/cgRosterArtwork.generated.json",
        "revisions": "src/game/cg/cgRevisions.generated.json",
    }
    for key, filename in expected_inputs.items():
        pin = manifest["inputs"][key]
        assert pin["path"] == filename
        assert pin["sha256"] == hashlib.sha256((ROOT / filename).read_bytes()).hexdigest(), f"Stale global manifest: {key}"
    subjects = {subject["characterId"]: subject for subject in plan["subjects"]}
    assert len(manifest["subjects"]) == len(subjects)
    assert {subject["characterId"] for subject in manifest["subjects"]} == set(subjects)
    expected_slots = {f"CG{i:02d}" for i in range(1, 16)}
    dispositions = Counter()
    for subject in manifest["subjects"]:
        jobs = subject["jobs"]
        assert len(jobs) == 15 and {job["slot"] for job in jobs} == expected_slots
        for job in jobs:
            assert job["production"] == subjects[subject["characterId"]]["matrix"][job["slot"]], "A global instruction changed the real per-slot production facts"
            assert job["fidelityCertifiedByManifest"] is False
            dispositions[job["queueStatus"]] += 1
    counts = manifest["summary"]
    assert counts == batch["summary"]
    assert counts["cgSlots"] == len(subjects) * 15
    assert counts["actualDeliveredCg"] == plan["summary"]["generated"]
    assert counts["documentedNotApplicable"] == plan["summary"]["matrixNotApplicableTotal"]
    assert counts["remainingSlots"] == plan["summary"]["totalRequestedSlotsNotProduced"]
    assert counts["queueStatuses"] == dict(dispositions)
    assert counts["imageCallsMadeByThisBuilder"] == counts["imagesInstalledByThisBuilder"] == counts["automaticPerfectFidelityApprovals"] == 0
    assert counts["currentProviderRefusedOrCorrectionRefusedJobs"] == 20
    assert counts["currentAtlasRefusedSubjects"] == 7 and counts["newAtlasRefusedSubjects"] == 1
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    rel_id = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    with ZipFile(book) as archive:
        sheet = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        links = {node.get("ref"): node.get(rel_id) for node in sheet.findall("s:hyperlinks/s:hyperlink", ns)}
        targets = {node.get("Id"): node.get("Target") for node in ET.fromstring(archive.read("xl/worksheets/_rels/sheet1.xml.rels"))}
        for cell, expected in [("B41", prompt_path), ("B42", manifest_path)]:
            target = unquote(targets[links[cell]].split("#", 1)[0])
            assert (book.parent / target).resolve() == expected.resolve(), f"Broken global entry point: {cell}"
    return {"manifest": batch["manifestPath"], "manifestSha256": batch["manifestSha256"],
            "masterPrompt": batch["masterPromptPath"], "masterPromptSha256": batch["masterPromptSha256"],
            "subjects": len(subjects), "slots": sum(dispositions.values()),
            "actualDeliveredCg": counts["actualDeliveredCg"], "currentCgRefusals": 20,
            "currentAtlasRefusals": 7, "newAtlasRefusals": 1,
            "homeEntryPointLinksVerified": 2, "instructionsMistakenForGeneratedImages": 0}


def audit(book):
    plan = read_json("docs/cg/roster-production-2026-10-18.json")
    bible = read_json("docs/cg/roster-fidelity-bible-2026-10-18.json")
    baseline = read_json("public/cg/cg-manifest.json")["entries"]
    roster = read_json("src/game/cg/cgRosterArtwork.generated.json")["entries"]
    revisions = read_json("src/game/cg/cgRevisions.generated.json")["revisions"]
    latest = {entry["cgId"]: entry for entry in revisions}
    subjects = {subject["characterId"] for subject in plan["subjects"]}
    plan_by_id = {subject["characterId"]: subject for subject in plan["subjects"]}
    partial_wave_states = {}
    for character_id, name in [("stitch_626", "stitch"), ("raven_tt", "raven"), ("starfire_tt", "starfire"), ("jack_sparrow_potc", "jack-sparrow")]:
        receipt_file = f"docs/cg/{name}-wave3-2026-10-20.json"
        receipt = read_json(receipt_file)
        for blocked, values in [(True, receipt.get("blockages", []) + receipt.get("blockedSlots", []) + receipt.get("toolFailures", [])), (False, receipt.get("unattemptedSlots", []))]:
            for value in values:
                raw = value.get("slot") if isinstance(value, dict) else value
                slot = f"CG{int(str(raw).removeprefix('CG')):02d}"
                categories = set()
                if isinstance(value, dict):
                    categories.update(value.get("categories", []))
                    if value.get("category"):
                        categories.add(value["category"])
                partial_wave_states[(character_id, slot)] = (blocked, categories, receipt_file)
    assert subjects == {subject["characterId"] for subject in bible["subjects"]}, "Bible/roster coverage differs"
    expected_media = {}
    for entry in baseline + roster:
        slot = next(f"CG{index:02d}" for index, types in enumerate(SLOTS, 1) if entry["type"] in types)
        key = entry["characterId"], slot
        assert key not in expected_media, f"Conflicting CG alternatives: {key}"
        current = latest.get(entry["id"])
        if current:
            paths = current["imagePath"], current["sourcePath"]
        else:
            paths = entry["assets"]["master"]["path"], entry["assets"]["source"]["path"]
        for name in paths:
            assert (ROOT / "public" / name.lstrip("/")).is_file(), f"Missing actual media: {name}"
        expected_media[key] = paths

    workbook = load_workbook(book, read_only=True, data_only=False)
    expected_sheets = {"Accueil", "Roster", "Contrôles", "Matrice CG", "CG livrées", "Variantes", "Univers", "Sources", "Comparaisons", "Variantes sprites", "Règles CG", "Traçabilité", "Listes"}
    assert expected_sheets <= set(workbook.sheetnames), "Missing review sheets"
    roster_rows = list(rows(workbook["Roster"]))
    assert len(roster_rows) == len(subjects)
    assert {row["ID personnage"] for row in roster_rows} == subjects, "Roster has omissions or duplicates"
    assert len({row["Univers"] for row in roster_rows}) == plan["summary"]["universes"]
    assert all(row["Incarnation / continuité"] for row in roster_rows), "A missing identity must be explicit"
    bible_by_id = {subject["characterId"]: subject for subject in bible["subjects"]}
    selected_sprite_checks = 0
    for row in roster_rows:
        selected = next((visual.get("path") for visual in bible_by_id[row["ID personnage"]].get("localVisuals", [])
                         if visual.get("selectedInRuntime") and visual.get("exists")), None)
        if selected:
            assert row["Fichier visuel"] == selected, f"Excel shows an older sprite: {row['ID personnage']}"
            selected_sprite_checks += 1
    actual_keys = set()
    actual_images = {}
    not_applicable = 0
    partial_state_checks = {"refused": 0, "notAttempted": 0}
    mistral_state_checks = {"refused": 0, "notAttempted": 0}
    for row in rows(workbook["Matrice CG"]):
        key = row["ID personnage"], row["CG"]
        assert key[0] in subjects and key[1] in {f"CG{i:02d}" for i in range(1, 16)}
        assert key not in actual_keys, f"Duplicate matrix row: {key}"
        actual_keys.add(key)
        expected = expected_media.get(key)
        image = row["Image actuelle"]
        source = row["PNG natif / source"]
        if expected:
            assert (image, source) == expected, f"Excel does not show the latest registered media: {key}"
            actual_images[key] = image
        else:
            assert not image and not source, f"A planned CG is presented as delivered: {key}"
        assert row["État réel de production"], f"Missing production state: {key}"
        if key in partial_wave_states and not expected:
            blocked, categories, receipt_file = partial_wave_states[key]
            state = str(row["État réel de production"])
            limits = str(row["Limites connues"] or "")
            assert not state.startswith("N/A"), f"A provider refusal/non-attempt became N/A: {key}"
            assert receipt_file in limits, f"Missing actual generation receipt in Excel: {key}"
            assert row["Pose distincte prévue / livrée"] == plan_by_id[key[0]]["matrix"][key[1]]["poseBrief"], f"Authored pose lost in Excel: {key}"
            if blocked:
                assert "refus du service" in state.lower(), f"An actual provider refusal disappeared: {key}"
                assert all(category in limits for category in categories), f"Communicated refusal category lost: {key}"
                partial_state_checks["refused"] += 1
            else:
                assert "refus" not in state.lower() and "pas encore été tenté" in limits, f"An unattempted slot became a refusal: {key}"
                partial_state_checks["notAttempted"] += 1
        if str(row["État réel de production"]).startswith("N/A"):
            not_applicable += 1
        if not str(row["État réel de production"]).startswith("N/A"):
            assert row["Pose distincte prévue / livrée"] and row["Thème / contrainte"], f"Missing pose/theme: {key}"
        if key[0] == "mistral_mgr":
            state = str(row["État réel de production"])
            limits = str(row["Limites connues"] or "")
            if key[1] in {"CG01", "CG04", "CG09", "CG10"}:
                assert any(label in state.lower() for label in ("refus du service", "refusée par le service")), f"A real Mistral refusal was lost: {key}"
                assert "sexual" in limits, f"The communicated category was lost: {key}"
                if key[1] == "CG04":
                    assert "docs/cg/mistral-wave5-2026-10-22.json" in limits, "An empty-delivery CG receipt still documents an actual refusal"
                mistral_state_checks["refused"] += 1
            elif key[1] in {"CG05", "CG06", "CG07", "CG08", "CG11", "CG12", "CG13", "CG14", "CG15"}:
                assert "refus" not in state.lower() and "pas encore été tenté" in limits, f"A non-attempt became a refusal: {key}"
                mistral_state_checks["notAttempted"] += 1
    assert len(actual_keys) == len(subjects) * 15
    assert actual_images.keys() == expected_media.keys()
    assert not_applicable == plan["summary"]["matrixNotApplicableTotal"]
    assert len(list(rows(workbook["CG livrées"]))) == len(expected_media)
    assert len(list(rows(workbook["Variantes"]))) == len(baseline) + len(roster) + len(revisions)
    controls = list(rows(workbook["Contrôles"]))
    assert {row["ID personnage"] for row in controls} == subjects
    assert all(isinstance(row["Axes conformes"], str) and row["Axes conformes"].startswith("=") for row in controls), "Per-character checklist formula missing"
    scoped_reference_checks = {
        "danteEquipmentOnly": 0, "danteEnvironmentOnly": 0,
        "sonicCostumeOnly": 0, "raymanEnvironmentOnly": 0,
        "excludedNerisCopyRows": 0, "excludedNyraHistoricalComparisonRows": 0,
        "nyraInternalEquipmentReuse": 0, "nyraInternalEnvironmentReuse": 0,
        "jaxInternalEquipmentReuse": 0, "jaxInternalEnvironmentReuse": 0,
        "maraInternalEnvironmentOnly": 0,
    }
    v2_portrait_paths = {
        "elara_voss": "aether_crown", "nyra_vale": "neon_requiem",
        "jax_rook": "neon_requiem", "seraphine_meridius": "meridian_clockwork",
        "mara_veylan": "iron_tempest", "celeste_ray": "helios_56",
        "nefra_balance": "kemet_devoured_sun",
        "brann_crow": "crownless_sea", "lyra_hexwind": "aetherion_seven_laws",
        "mercy_ash": "infernum_nine_pits", "mara_ell": "drowned_testament",
        "orin_vale": "aetherion_seven_laws",
    }
    v2_original_identities = set()
    for row in rows(workbook["Comparaisons"]):
        source_file = str(row["Fichier source local"] or "").lower()
        usage = str(row["Usage de la référence / CG"] or "")
        axes = str(row["Axes à comparer"] or "")
        character_id = row["ID personnage"]
        if character_id == "neris_vale" and source_file.endswith("/current-atlas-source-copy.png"):
            scoped_reference_checks["excludedNerisCopyRows"] += 1
        if character_id == "nyra_vale" and source_file.endswith("/selected-runtime-atlas-comparison.png"):
            scoped_reference_checks["excludedNyraHistoricalComparisonRows"] += 1
        if character_id in v2_portrait_paths and usage == "CG01 — identité / incarnation":
            assert "Autorité canon interne documentée : portrait original" in str(row["Statut de la source"] or ""), "The selected V2 portrait is internal project authority, never an external franchise reference"
            selected = ROOT / f"public/images/oc-worlds/v2/{v2_portrait_paths[character_id]}/heroes/{character_id}.png"
            actual_reference = Path(row["Fichier source local"])
            assert hashlib.sha256(actual_reference.read_bytes()).digest() == hashlib.sha256(selected.read_bytes()).digest(), f"The identity thumbnail must be the actual selected portrait bitmap: {character_id}"
            v2_original_identities.add(character_id)
        if character_id == "jax_rook" and usage in {"CG02 — équipement uniquement", "CG03 — décor / univers uniquement"}:
            status = str(row["Statut de la source"] or "")
            assert "Autorité canon interne documentée : portrait original" in status
            assert "Réutilisation du même bitmap" in status
            selected = ROOT / "public/images/oc-worlds/v2/neon_requiem/heroes/jax_rook.png"
            assert hashlib.sha256(Path(row["Fichier source local"]).read_bytes()).digest() == hashlib.sha256(selected.read_bytes()).digest()
            if usage.startswith("CG02"):
                assert axes.startswith("Équipement uniquement")
                scoped_reference_checks["jaxInternalEquipmentReuse"] += 1
            else:
                assert axes.startswith("Décor / univers uniquement")
                scoped_reference_checks["jaxInternalEnvironmentReuse"] += 1
        if character_id == "mara_veylan" and source_file.endswith("/iron_tempest/backdrop.png"):
            assert usage == "CG03 — décor / univers uniquement" and axes.startswith("Décor / univers uniquement")
            assert "Autorité canon interne documentée : décor original" in str(row["Statut de la source"] or "")
            scoped_reference_checks["maraInternalEnvironmentOnly"] += 1
        if character_id == "nyra_vale" and usage in {"CG02 — équipement uniquement", "CG03 — décor / univers uniquement"}:
            assert "Autorité canon interne documentée : portrait original" in str(row["Statut de la source"] or ""), "Reusing Nyra's original portrait for equipment/environment must keep its internal authority"
            reference_path = Path(row["Fichier source local"])
            assert reference_path.is_file()
            assert hashlib.sha256(reference_path.read_bytes()).hexdigest() == "8f14ea6f6daa0bd8dcb8dda860ac72549356fb520387ab970351781f7ce5255f", "Scoped reuse must point to the same qualified original Nyra portrait"
            if usage.startswith("CG02"):
                assert axes.startswith("Équipement uniquement")
                scoped_reference_checks["nyraInternalEquipmentReuse"] += 1
            else:
                assert axes.startswith("Décor / univers uniquement")
                scoped_reference_checks["nyraInternalEnvironmentReuse"] += 1
        if row["ID personnage"] == "devil_may_cry_dante" and "weapon-dmc5_ebony_and_ivory" in source_file:
            assert usage == "CG02 — équipement uniquement" and axes.startswith("Équipement uniquement"), "Dante's weapon reference must never approve character identity"
            scoped_reference_checks["danteEquipmentOnly"] += 1
        if row["ID personnage"] == "devil_may_cry_dante" and "dmc5-steam-screen-" in source_file:
            assert usage == "CG03 — décor / univers uniquement" and axes.startswith("Décor / univers uniquement"), "Nero/V environment screenshots must never approve Dante's identity"
            scoped_reference_checks["danteEnvironmentOnly"] += 1
        if row["ID personnage"] == "sonic_hedgehog" and "riders" in source_file:
            assert usage == "CG13 — costume croisé What If seulement" and axes.startswith("Tenue croisée uniquement"), "Riders must not replace Modern Sonic's identity source"
            scoped_reference_checks["sonicCostumeOnly"] += 1
        if character_id == "rayman_rayman" and source_file.endswith("/steam-screenshot-01.jpg"):
            assert usage == "CG03 — décor / univers uniquement" and axes.startswith("Décor / univers uniquement"), "Rayman's empty woodland concept art cannot qualify character identity"
            scoped_reference_checks["raymanEnvironmentOnly"] += 1
    assert scoped_reference_checks["danteEquipmentOnly"] >= 1
    assert scoped_reference_checks["danteEnvironmentOnly"] == 5
    assert scoped_reference_checks["sonicCostumeOnly"] >= 1
    assert scoped_reference_checks["raymanEnvironmentOnly"] >= 1
    assert scoped_reference_checks["excludedNerisCopyRows"] == 0, "A byte-identical Neris source copy became an independent identity reference"
    assert scoped_reference_checks["excludedNyraHistoricalComparisonRows"] == 0, "Nyra's divergent historical atlas became an independent V2 identity reference"
    assert scoped_reference_checks["nyraInternalEquipmentReuse"] == 1
    assert scoped_reference_checks["nyraInternalEnvironmentReuse"] == 1
    assert scoped_reference_checks["jaxInternalEquipmentReuse"] == 1
    assert scoped_reference_checks["jaxInternalEnvironmentReuse"] == 1
    assert scoped_reference_checks["maraInternalEnvironmentOnly"] == 1
    assert v2_original_identities == set(v2_portrait_paths)
    assert mistral_state_checks == {"refused": 4, "notAttempted": 9}
    workbook.close()

    embedded_images = []
    local_links = 0
    with ZipFile(book) as archive:
        assert archive.testzip() is None, "Corrupt Excel archive"
        embedded_images = [name for name in archive.namelist() if name.startswith("xl/media/")]
        assert len(embedded_images) >= len(subjects) + len(expected_media), "Thumbnails are not embedded"
        for name in archive.namelist():
            if not name.endswith(".rels"):
                continue
            for relationship in ET.fromstring(archive.read(name)):
                if not relationship.attrib.get("Type", "").endswith("/hyperlink"):
                    continue
                target = unquote(relationship.attrib["Target"])
                if target.startswith(("http://", "https://", "#")):
                    continue
                local = (book.parent / target.split("#", 1)[0]).resolve()
                assert local.is_file(), f"Broken local Excel link: {target}"
                local_links += 1
    global_batch_checks = audit_global_batch(book, plan)
    return {
        "schemaVersion": 1,
        "workbook": str(book.relative_to(ROOT)),
        "sha256": hashlib.sha256(book.read_bytes()).hexdigest(),
        "subjects": len(subjects), "universes": plan["summary"]["universes"],
        "matrixRows": len(actual_keys), "registeredCurrentCg": len(expected_media),
        "selectedSpritesMatched": selected_sprite_checks,
        "retainedVersions": len(baseline) + len(roster) + len(revisions),
        "notApplicable": not_applicable, "embeddedImages": len(embedded_images),
        "verifiedLocalHyperlinks": local_links,
        "scopedReferenceUsageChecks": scoped_reference_checks,
        "verifiedInternalV2PortraitSubjects": sorted(v2_original_identities),
        "partialWaveRefusalAndNonAttemptChecks": partial_state_checks,
        "mistralRefusalAndNonAttemptChecks": mistral_state_checks,
        "globalBatchChecks": global_batch_checks,
        "allCurrentMediaMatchRuntime": True, "plannedImagesPresentedAsDelivered": 0,
        "certifiesVisualResemblance": False,
    }


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--write", action="store_true")
    options = parser.parse_args()
    result = audit(BOOK)
    if options.write:
        BOOK.with_suffix(".validation.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
